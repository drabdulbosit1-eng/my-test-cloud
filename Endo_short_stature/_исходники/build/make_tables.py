"""Таблица источников для сдачи: литература, утверждения, понятия, заголовки.

python3 make_tables.py <папка _исходники> <выход.xlsx>
"""
import glob
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

src, out = sys.argv[1], sys.argv[2]

slides = []
for f in sorted(glob.glob(os.path.join(src, 'текст', '*.json'))):
    slides += json.load(open(f, encoding='utf-8'))
slides.sort(key=lambda s: s['n'])

wb = Workbook()
head_font = Font(bold=True, color='FFFFFF')
head_fill = PatternFill('solid', fgColor='1B2E3C')
wrap = Alignment(wrap_text=True, vertical='top')


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill
        c.alignment = wrap
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = 'A2'


# 1. Литература (слайд 43)
refs = next(s for s in slides if s['n'] == 43)['body']
ws = wb.active
ws.title = 'Литература'
rows = []
for r in refs:
    m = re.match(r'^(\d+)\.\s*(.*)$', r)
    rows.append([int(m.group(1)) if m else None, m.group(2) if m else r])
sheet(ws, ['№', 'Источник'], rows, [6, 140])

# 2. Утверждения: все строки проверки
ws = wb.create_sheet('Утверждения')
rows = []
for f in sorted(glob.glob(os.path.join(src, 'проверка', '*.md'))):
    part = os.path.basename(f).replace('.md', '')
    for line in open(f, encoding='utf-8'):
        if not line.startswith('| ') or line.startswith('| ID') or line.startswith('|---'):
            continue
        c = [x.strip() for x in line.strip().strip('|').split('|')]
        if len(c) < 8 or c[0] in ('прим.',):
            continue
        rows.append([part] + c[:8])
sheet(ws, ['Файл проверки', 'ID', 'Утверждение', 'Источник', 'PMID', 'DOI', 'Уровень доказательности', 'Статус',
           'Цитата'], rows, [16, 9, 60, 40, 11, 26, 30, 24, 60])

# 3. Понятия — слайд
ws = wb.create_sheet('Понятия по слайдам')
rows = []
for s in slides:
    for c in s.get('concepts_home') or []:
        rows.append([c, s['n']])
rows.sort(key=lambda r: r[1])
sheet(ws, ['Понятие', 'Домашний слайд'], rows, [55, 16])

# 4. Заголовки подряд и источник каждого слайда
ws = wb.create_sheet('Слайды')
rows = [[s['n'], s['title'], s.get('evidence') or '', s.get('source') or '',
         len((s.get('notes') or '').split())] for s in slides]
sheet(ws, ['№', 'Заголовок', 'Уровень доказательности', 'Источник', 'Слов в заметках'], rows, [5, 70, 45, 60, 10])

wb.save(out)
print('saved', out, len(refs), 'refs', len(slides), 'slides')
