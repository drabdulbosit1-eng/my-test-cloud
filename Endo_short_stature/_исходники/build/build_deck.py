"""Сборка лекции из JSON-файлов текста.

python3 build_deck.py <папка_с_json> <выход.pptx>
"""
import copy
import glob
import json
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
from style import (BODY_BOTTOM, BODY_Y, C, CONTENT_W, EVID_Y, FONT_BODY, FONT_HEAD, MX, PARTS, RULER_Y,
                   SIZE_BODY, SIZE_EVID, SIZE_SRC, SIZE_TABLE, SIZE_TITLE, SLIDE_H, SLIDE_W, SRC_H, SRC_TOP, SRC_Y,
                   TITLE_H, TITLE_Y, TOTAL)
import growth
import who
from figures import FIGS

NB = ' '
UNITS = (r'мг|мкг|мл|кг|см|мм|м²|лет|года|год|мес|нед|дн|сут|ч|нмоль/л|пмоль/л|нг/мл|пг/мл|мЕд/л|МЕ/л|'
         r'мкг/дл|мг/л|%|Гр|кДа|раз|случа|детей|ребён|исследован|пациент|мальчик|девоч|человек|из|СО')


def typo(s):
    """Неразрывные пробелы: число и единица, знаки сравнения, n = 24."""
    if not s:
        return s
    s = re.sub(r'(\d) (?=(' + UNITS + r'))', r'\1' + NB, s)
    s = re.sub(r'([<>≤≥=±−]) (?=[\d−])', r'\1' + NB, s)
    s = re.sub(r'(\bn|\bр|\bp) ([=<>])', r'\1' + NB + r'\2', s)
    s = re.sub(r'(\d) (–|−) (\d)', r'\1\2\3', s)
    s = re.sub(r'(\d)–(\d)', '\\1–\u2060\\2', s)   # не разрывать диапазон «3–6»
    s = s.replace('св. Т4', 'св.' + NB + 'Т4')
    s = re.sub(r'(?<=[А-ЯA-Z])-(?=\d)', '-\u2060', s)       # ИФР-1, ИФРСБ-3 без переноса
    return s


def rgb(h):
    return RGBColor.from_string(h)


def add_runs(par, text, size, color, font=FONT_BODY, bold=False):
    """Текст с курсивом для *генов*."""
    parts = re.split(r'(\*[^*]+\*)', typo(text))
    for part in parts:
        if not part:
            continue
        r = par.add_run()
        it = part.startswith('*') and part.endswith('*')
        r.text = part[1:-1] if it else part
        f = r.font
        f.size = size
        f.name = font
        f.bold = bold
        f.italic = it
        f.color.rgb = rgb(color)


def textbox(shapes, x, y, w, h, anchor=MSO_ANCHOR.TOP, name=None):
    tb = shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if name:
        tb.name = name
    return tb, tf


def set_bullet(par, level_indent=Inches(0.28), char='–'):
    pPr = par._p.get_or_add_pPr()
    pPr.set('marL', str(int(level_indent)))
    pPr.set('indent', str(-int(level_indent)))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn('a:buClr'))
    srgb = etree.SubElement(buClr, qn('a:srgbClr'))
    srgb.set('val', C['height'])
    bu = etree.SubElement(pPr, qn('a:buChar'))
    bu.set('char', char)


def space_after(par, pt):
    par.space_after = Pt(pt)


def title_size(text):
    n = len(text)
    return Pt(28) if n <= 80 else (Pt(25) if n <= 90 else Pt(23))


def title(slide, text, dark=False):
    tb, tf = textbox(slide.shapes, MX, TITLE_Y, CONTENT_W, TITLE_H, anchor=MSO_ANCHOR.TOP, name='Заголовок')
    p = tf.paragraphs[0]
    p.line_spacing = 0.95
    add_runs(p, text, title_size(text), C['onink'] if dark else C['ink'], font=FONT_HEAD, bold=True)


def bullets(slide, items, x, y, w, h, size=SIZE_BODY, color=None, gap=10):
    tb, tf = textbox(slide.shapes, x, y, w, h, name='Текст')
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        set_bullet(p)
        space_after(p, gap)
        p.line_spacing = 1.0
        add_runs(p, it, size, color or C['text'])
    return tb


def footer(slide, n, evidence, source, dark=False):
    """Только строка источника внизу. По правке автора нет ни строки доказательности,
    ни шкалы разделов, ни номера слайда; уровень доказательности — в таблице источников."""
    col = C['onink2'] if dark else C['muted']
    if source:
        tb, tf = textbox(slide.shapes, MX, SRC_TOP, CONTENT_W, SRC_H, anchor=MSO_ANCHOR.BOTTOM, name='Источник')
        ln_ = len(source)
        sz = Pt(10) if ln_ <= 260 else (Pt(9) if ln_ <= 380 else Pt(8))
        add_runs(tf.paragraphs[0], source, sz, col)


def ruler(slide, n, dark=False):
    """Шкала-ростомер: где слушатель находится в лекции."""
    sh = slide.shapes
    x0, x1 = int(MX), int(SLIDE_W - MX)
    y = int(RULER_Y + Inches(0.05))
    base_col = C['onink2'] if dark else C['rule']
    fill_col = C['onink'] if dark else C['height']

    def rect(x, yy, w, h, col, name):
        r = sh.add_shape(1, int(x), int(yy), max(int(w), 1), int(h))
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(col)
        r.line.fill.background()
        r.shadow.inherit = False
        r.name = name
        return r

    step = (x1 - x0) / TOTAL
    rect(x0, y, x1 - x0, Inches(0.03), base_col, 'Ростомер: шкала')
    rect(x0, y - Inches(0.015), step * n, Inches(0.06), fill_col, 'Ростомер: пройдено')
    for i in range(TOTAL + 1):
        major = any(i == a - 1 for a, b, _ in PARTS) or i == TOTAL
        hgt = Inches(0.16) if major else Inches(0.07)
        rect(x0 + step * i - Emu(4000), y - hgt, Emu(8000), hgt, base_col, 'Ростомер: деление')
    for a, b, name in PARTS:
        cx = x0 + step * (a - 1)
        w = step * (b - a + 1)
        cur = a <= n <= b
        tb, tf = textbox(sh, cx, y + Inches(0.07), w, Inches(0.22), name='Ростомер: раздел')
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        col = (C['onink'] if dark else C['ink']) if cur else (C['onink2'] if dark else C['muted'])
        add_runs(p, name, Pt(10), col, bold=cur)
    tb, tf = textbox(sh, x1 - Inches(0.6), y - Inches(0.36), Inches(0.6), Inches(0.24), name='Номер слайда')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_runs(p, f'{n}', Pt(11), C['onink2'] if dark else C['muted'])


def background(slide, col):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(col)


def table(slide, spec, x, y, w, h, size=SIZE_TABLE, col_w=None):
    header, rows = spec['header'], spec['rows']
    nr, nc = len(rows) + 1, len(header)
    gt = slide.shapes.add_table(nr, nc, int(x), int(y), int(w), int(h))
    gt.name = 'Таблица'
    tbl = gt.table
    # без встроенного стиля
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '1')
    tblPr.set('bandRow', '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is None:
        style = etree.SubElement(tblPr, qn('a:tableStyleId'))
    style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # No Style, Table Grid
    if col_w:
        tot = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = int(w * cw / tot)
    rh = est_rows(spec, w, size, col_w)
    for i, hh in enumerate(rh):
        tbl.rows[i].height = int(Inches(hh))
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            txt = header[c] if r == 0 else (rows[r - 1][c] if c < len(rows[r - 1]) else '')
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            for run in list(p.runs):
                run._r.getparent().remove(run._r)
            add_runs(p, str(txt), size, C['onink'] if r == 0 else C['text'], bold=(r == 0))
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(C['ink'] if r == 0 else (C['paper2'] if r % 2 == 0 else C['paper']))
            # тонкие линии
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = tcPr.find(qn(side))
                if ln is None:
                    ln = etree.SubElement(tcPr, qn(side))
                ln.set('w', '6350')
                for ch in list(ln):
                    ln.remove(ch)
                sf = etree.SubElement(ln, qn('a:solidFill'))
                cl = etree.SubElement(sf, qn('a:srgbClr'))
                cl.set('val', C['rule'])
    return gt


def est_rows(tbl, w, size, col_w=None):
    """Оценка высоты каждой строки таблицы после переноса (дюймы)."""
    import math
    header, rows = tbl['header'], tbl['rows']
    nc = len(header)
    cw = col_w or [1] * nc
    tot = sum(cw)
    widths = [w / 914400 * c / tot - 0.18 for c in cw]
    char_w = size.pt / 72 * 0.53
    line_h = size.pt / 72 * 1.25
    out = []
    for r in [header] + rows:
        lines = 1
        for c in range(nc):
            txt = str(r[c]) if c < len(r) else ''
            per = max(1, int(widths[c] / char_w))
            lines = max(lines, math.ceil(max(1, len(txt)) / per))
        out.append(lines * line_h + 0.14)
    return out


def est_table_h(tbl, w, size, col_w=None):
    return Inches(sum(est_rows(tbl, w, size, col_w)))


def table_size(tbl):
    chars = sum(len(str(c)) for r in tbl['rows'] for c in r)
    if chars <= 330:
        return Pt(18)
    return Pt(16) if chars <= 420 else (Pt(14) if chars <= 600 else Pt(13))


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = typo(text).replace('*', '')


def caption(slide, x, y, w, text='схема, а не данные конкретного пациента'):
    tb, tf = textbox(slide.shapes, x, y, w, Inches(0.24), name='Подпись к схеме')
    add_runs(tf.paragraphs[0], text, Pt(11), C['muted'])
    tf.paragraphs[0].runs[0].font.italic = True


def legend(slide, x, y):
    items_ = ((C['height'], 'пример: рост'), (C['weight'], 'пример: ИМТ'), (C['band'], 'ВОЗ: 3–97-й центиль')) \
        if who.available() else ((C['height'], 'пример: рост'), (C['weight'], 'пример: масса'),
                                 (C['band'], 'норма: 3–97-й центиль'))
    for i, (col, txt) in enumerate(items_):
        r = slide.shapes.add_shape(1, int(x + Inches(2.2) * i), int(y + Inches(0.07)), int(Inches(0.3)),
                                   int(Inches(0.1)))
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(col)
        r.line.fill.background()
        r.shadow.inherit = False
        r.name = 'Легенда'
        tb, tf = textbox(slide.shapes, x + Inches(2.2) * i + Inches(0.38), y, Inches(1.8), Inches(0.24),
                         name='Легенда')
        add_runs(tf.paragraphs[0], txt, Pt(11), C['muted'])


# ---------- типы слайдов ----------

LOGO = os.path.join(os.path.dirname(__file__), '..', 'версия_автора', 'рисунки', 'логотип.png')


def s_title(prs, d):
    """Титул по образцу автора: логотип, название, имя и специальность."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, C['paper'])
    sh = s.shapes
    pic = sh.add_picture(LOGO, int(MX), int(Inches(0.9)), width=int(Inches(5.2)))
    pic.name = 'Логотип'
    tb, tf = textbox(sh, MX, Inches(3.35), Inches(10.5), Inches(1.75), anchor=MSO_ANCHOR.BOTTOM, name='Название')
    p = tf.paragraphs[0]
    p.line_spacing = 0.95
    add_runs(p, d['title'], Pt(38), C['ink'], font=FONT_HEAD, bold=True)
    body = d.get('body') or []
    tb, tf = textbox(sh, MX, Inches(5.45), Inches(10.5), Inches(1.0), name='Автор')
    for i, line in enumerate(body):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        add_runs(p, line, Pt(20) if i == 0 else Pt(16), C['text'] if i == 0 else C['muted'], bold=(i == 0))
    notes(s, d.get('notes'))
    return s


def s_thanks(prs, d):
    """Последний слайд по образцу автора."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, C['paper'])
    sh = s.shapes
    tb, tf = textbox(sh, MX, Inches(2.0), CONTENT_W, Inches(1.0), anchor=MSO_ANCHOR.BOTTOM, name='Заголовок')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, d['title'], Pt(40), C['ink'], font=FONT_HEAD, bold=True)
    body = d.get('body') or []
    tb, tf = textbox(sh, Inches(2.4), Inches(3.9), Inches(4.2), Inches(1.2), anchor=MSO_ANCHOR.MIDDLE, name='Контакт')
    for i, line in enumerate(body):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        add_runs(p, line, Pt(26) if i == 0 else Pt(18), C['text'] if i == 0 else C['muted'])
    pic = sh.add_picture(LOGO, int(Inches(7.1)), int(Inches(3.75)), width=int(Inches(3.9)))
    pic.name = 'Логотип'
    notes(s, d.get('notes'))
    return s


def s_dark(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, C['ink'])
    title(s, d['title'], dark=True)
    items = d.get('body') or []
    bullets(s, items, MX, BODY_Y, CONTENT_W, BODY_BOTTOM - BODY_Y, size=Pt(20), color=C['onink'], gap=12)
    footer(s, d['n'], d.get('evidence'), d.get('source'), dark=True)
    notes(s, d.get('notes'))
    return s


def s_content(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, C['paper'])
    title(s, d['title'])
    n = d['n']
    fig = d.get('figure') or FIGS.get(n)
    items = d.get('body') or []
    tbl = d.get('table')
    if tbl and not tbl.get('rows'):
        tbl = None
    top, bottom = BODY_Y, BODY_BOTTOM
    if fig and fig['type'] == 'growth' and who.available():
        fw = Inches(8.2) if items else CONTENT_W
        who.chart(s.shapes, MX, top + Inches(0.3), fw, bottom - top - Inches(0.95), fig['sex'], fig['t0'], fig['t1'],
                  fig['hz'], fig['wz'], fig.get('events', ()))
        legend(s, MX, bottom - Inches(0.3))
        if items:
            bullets(s, items, MX + fw + Inches(0.4), top + Inches(0.3), CONTENT_W - fw - Inches(0.4),
                    bottom - top - Inches(0.3), size=Pt(18))
    elif fig and fig['type'] == 'two_growth' and who.available():
        half = (CONTENT_W - Inches(0.5)) / 2
        h = bottom - top - Inches(0.95) - (Inches(0.36) * len(items) + Inches(0.15) if items else 0)
        for i, g in enumerate(fig['charts']):
            who.chart(s.shapes, MX + i * (half + Inches(0.5)), top + Inches(0.3), half, h, g['sex'], g['t0'], g['t1'],
                      g['hz'], g['wz'], g.get('events', ()), title=g['title'])
        legend(s, MX, top + h + Inches(0.6))
        if items:
            bullets(s, items, MX, top + h + Inches(0.95), CONTENT_W, Inches(0.36) * len(items), size=Pt(17), gap=4)
    elif fig and fig['type'] == 'growth':
        fx, fw = MX, Inches(7.4)
        growth.growth_chart(s.shapes, fx, top, fw, bottom - top - Inches(0.85), fig['sex'], fig['t0'], fig['t1'],
                            fig['hz'], fig['wz'], fig.get('events', ()))
        legend(s, fx, bottom - Inches(0.52))
        bullets(s, items, MX + Inches(7.75), top + Inches(0.2), CONTENT_W - Inches(7.75), bottom - top - Inches(0.2),
                size=Pt(18))
    elif fig and fig['type'] == 'two_growth':
        half = (CONTENT_W - Inches(0.5)) / 2
        h = bottom - top - Inches(0.62) - (Inches(0.2) + Inches(0.36) * len(items) if items else 0)
        for i, g in enumerate(fig['charts']):
            growth.growth_chart(s.shapes, MX + i * (half + Inches(0.5)), top, half, h, g['sex'], g['t0'], g['t1'],
                                g['hz'], g['wz'], g.get('events', ()), title=g['title'])
        legend(s, MX, top + h + Inches(0.27))
        if items:
            bullets(s, items, MX, top + h + Inches(0.62), CONTENT_W, Inches(0.36) * len(items), size=Pt(17), gap=4)
    elif fig and fig['type'] == 'thumbs':
        cols, rows_ = fig.get('cols', 3), fig.get('rows', 2)
        gw, gh = Inches(0.35), Inches(0.25)
        area_w = CONTENT_W - (Inches(3.9) if items else 0)
        tw = (area_w - gw * (cols - 1)) / cols
        th = (bottom - top - Inches(0.35) - gh * (rows_ - 1)) / rows_
        for i, t in enumerate(fig['items']):
            r, c = divmod(i, cols)
            growth.z_thumb(s.shapes, MX + c * (tw + gw), top + r * (th + gh), tw, th, t['title'], t['t0'], t['t1'],
                           t['hz'], t['wz'], t.get('note'))
        # легенда
        for i, (col, txt) in enumerate(((C['height'], 'рост'), (C['weight'], 'масса'),
                                        (C['band'], 'коридор нормы'))):
            rr = s.shapes.add_shape(1, int(MX + Inches(2.0) * i), int(bottom - Inches(0.22)), int(Inches(0.3)),
                                    int(Inches(0.1)))
            rr.fill.solid()
            rr.fill.fore_color.rgb = rgb(col)
            rr.line.fill.background()
            rr.shadow.inherit = False
            tb, tf = textbox(s.shapes, MX + Inches(2.0) * i + Inches(0.38), bottom - Inches(0.29), Inches(1.6),
                             Inches(0.24), name='Легенда')
            add_runs(tf.paragraphs[0], txt, Pt(11), C['muted'])
        if items:
            bullets(s, items, MX + area_w + Inches(0.4), top + Inches(0.2), CONTENT_W - area_w - Inches(0.4),
                    bottom - top - Inches(0.6), size=Pt(17))
    elif tbl and items:
        size = table_size(tbl)
        bh = Inches(0.34) * len(items) + Inches(0.1)
        th = est_table_h(tbl, CONTENT_W, size, d.get('col_w'))
        while th > bottom - top - bh - Inches(0.2) and size.pt > 12:
            size = Pt(size.pt - 1)
            th = est_table_h(tbl, CONTENT_W, size, d.get('col_w'))
        table(s, tbl, MX, top, CONTENT_W, th, size=size, col_w=d.get('col_w'))
        by = min(top + th + Inches(0.3), bottom - bh)
        bullets(s, items, MX, by, CONTENT_W, bottom - by, size=Pt(16), gap=3)
    elif tbl:
        size = table_size(tbl)
        th = est_table_h(tbl, CONTENT_W, size, d.get('col_w'))
        while th > bottom - top and size.pt > 12:
            size = Pt(size.pt - 1)
            th = est_table_h(tbl, CONTENT_W, size, d.get('col_w'))
        table(s, tbl, MX, top, CONTENT_W, min(th, bottom - top), size=size, col_w=d.get('col_w'))
    else:
        call = d.get('callout')
        if call:
            cw = Inches(3.6)
            tb, tf = textbox(s.shapes, MX, top + Inches(0.2), cw, Inches(1.4), anchor=MSO_ANCHOR.BOTTOM, name='Число')
            add_runs(tf.paragraphs[0], call['value'], Pt(54), C['height'], font=FONT_HEAD, bold=True)
            tb, tf = textbox(s.shapes, MX, top + Inches(1.7), cw, Inches(1.8), name='Подпись к числу')
            add_runs(tf.paragraphs[0], call['label'], Pt(16), C['muted'])
            bullets(s, items, MX + cw + Inches(0.5), top + Inches(0.2), CONTENT_W - cw - Inches(0.5),
                    bottom - top - Inches(0.2))
        else:
            bullets(s, items, MX, top + Inches(0.15), Inches(10.5), bottom - top - Inches(0.15))
    footer(s, n, d.get('evidence'), d.get('source'))
    notes(s, d.get('notes'))
    return s


def s_diagram(prs, d):
    """Слайд-схема механизма (diagrams.py)."""
    import diagrams
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, C['paper'])
    title(s, d['title'])
    diagrams.DRAW[d['diagram']['kind']](s, d)
    footer(s, d['n'], None, d.get('source'))
    notes(s, d.get('notes'))
    return s


def s_refs(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, C['paper'])
    title(s, d['title'])
    items = d.get('body') or []
    ncol = 3 if len(items) > 30 else 2
    size = Pt(8) if len(items) <= 100 else Pt(7)
    per = (len(items) + ncol - 1) // ncol
    gap = Inches(0.3)
    colw = (CONTENT_W - gap * (ncol - 1)) / ncol
    for i in range(ncol):
        chunk = items[i * per:(i + 1) * per]
        tb, tf = textbox(s.shapes, MX + i * (colw + gap), BODY_Y - Inches(0.35), colw,
                         RULER_Y - BODY_Y + Inches(0.2), name='Литература')
        for j, ref in enumerate(chunk):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(0.5)
            p.line_spacing = 0.95
            add_runs(p, ref, size, C['text'])
    ruler(s, d['n'])
    notes(s, d.get('notes'))
    return s


def load(folder):
    slides = []
    for f in sorted(glob.glob(os.path.join(folder, '*.json'))):
        with open(f, encoding='utf-8') as fh:
            slides.extend(json.load(fh))
    slides.sort(key=lambda d: d['n'])
    return slides


def build(folder, out):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for d in load(folder):
        lay = d.get('layout')
        if d['n'] == 1 or lay == 'title':
            s_title(prs, d)
        elif lay == 'thanks':
            s_thanks(prs, d)
        elif lay == 'diagram':
            s_diagram(prs, d)
        elif lay == 'references':
            s_refs(prs, d)
        elif lay == 'conclusions':
            s_dark(prs, d)
        else:
            s_content(prs, d)
    # без теней: обнулить ссылку на эффекты темы у всех фигур (LibreOffice иначе рисует тень)
    for sl in prs.slides:
        for ef in sl._element.iter(qn('a:effectRef')):
            ef.set('idx', '0')
    prs.core_properties.title = 'Эндокринные причины низкорослости у детей и подростков'
    prs.core_properties.language = 'ru-RU'
    prs.save(out)
    print('saved', out)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
