"""Схемы механизмов: зона роста, уровни оси, хромосома с геном SHOX.

Текст всех подписей лежит в JSON слайда (поле diagram), здесь только рисование.
Цвета из style.C, красный не используется.
"""
import os

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from style import BODY_BOTTOM, BODY_Y, C, CONTENT_W, FONT_HEAD, MX

CART = 'E4ECF0'      # хрящ
CELL = '9FB6C3'      # контур клетки
BONE = 'EDE4D0'      # кость
BONE_LN = 'C9B88E'

HERE = os.path.dirname(__file__)


def _rgb(h):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h)


def _box(sh, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, name=None):
    b = sh.add_shape(shape, int(x), int(y), int(w), int(h))
    if fill:
        b.fill.solid()
        b.fill.fore_color.rgb = _rgb(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = _rgb(line)
        b.line.width = Pt(lw)
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    if name:
        b.name = name
    return b


def _arrow(sh, x0, y0, x1, y1, col=None, w=1.5, name='Стрелка'):
    ln = sh.add_connector(MSO_CONNECTOR.STRAIGHT, int(x0), int(y0), int(x1), int(y1))
    ln.line.color.rgb = _rgb(col or C['axis'])
    ln.line.width = Pt(w)
    el = ln.line._get_or_add_ln()
    from lxml import etree
    from pptx.oxml.ns import qn
    tail = etree.SubElement(el, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    ln.name = name
    return ln


def _line(sh, x0, y0, x1, y1, col=None, w=0.75, name='Линия'):
    ln = sh.add_connector(MSO_CONNECTOR.STRAIGHT, int(x0), int(y0), int(x1), int(y1))
    ln.line.color.rgb = _rgb(col or C['rule'])
    ln.line.width = Pt(w)
    ln.name = name
    return ln


def _text(sh, x, y, w, h, runs, size=16, color=None, anchor=MSO_ANCHOR.TOP, align=None, name='Подпись',
          head=None, head_size=None, gap=4):
    """runs: строка или список абзацев. head — жирная первая строка."""
    from build_deck import add_runs, textbox
    tb, tf = textbox(sh, x, y, w, h, anchor=anchor, name=name)
    paras = [runs] if isinstance(runs, str) else list(runs)
    first = True
    if head:
        p = tf.paragraphs[0]
        first = False
        p.space_after = Pt(2)
        if align:
            p.alignment = align
        add_runs(p, head, Pt(head_size or size), color or C['ink'], bold=True)
    for t in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        if align:
            p.alignment = align
        add_runs(p, t, Pt(size), color or C['text'])
    return tb


def _image(sh, path, x, y, w, h, credit=None):
    """Картинка, вписанная в рамку; подпись автора и лицензии под ней."""
    from PIL import Image
    full = os.path.join(HERE, '..', path)
    if not os.path.exists(full):
        return None
    iw, ih = Image.open(full).size
    ch = Inches(0.42) if credit else 0
    s = min(w / iw, (h - ch) / ih)
    pw, ph = iw * s, ih * s
    px = x + (w - pw) / 2
    pic = sh.add_picture(full, int(px), int(y), int(pw), int(ph))
    pic.name = 'Изображение'
    if credit:
        _text(sh, x, y + ph + Inches(0.05), w, Inches(0.4), credit, size=9, color=C['muted'], name='Лицензия')
    return pic


# ---------- зона роста ----------

def growth_plate(slide, d):
    sh = slide.shapes
    g = d['diagram']
    top, bot = BODY_Y + Inches(0.05), BODY_BOTTOM - Inches(0.75)
    x0, cw = MX, Inches(3.0)
    img = g.get('image')
    zones = g['zones']                 # 4 зоны сверху вниз: покой, деление, созревание, кость
    frac = [0.18, 0.34, 0.26, 0.22]
    ys, y = [], top
    for f in frac:
        h = (bot - top) * f
        ys.append((y, h))
        y += h
    if img and os.path.exists(os.path.join(HERE, '..', img['path'])):
        _image(sh, img['path'], x0, top, cw, bot - top + Inches(0.5), img.get('credit'))
    else:
        import random
        rnd = random.Random(7)
        for i, (yy, hh) in enumerate(ys):
            _box(sh, x0, yy, cw, hh, fill=BONE if i == 3 else CART, name=f'Зона: {zones[i]["name"]}')
        # покой: редкие мелкие клетки
        yy, hh = ys[0]
        for k in range(14):
            cx = x0 + Inches(0.15) + rnd.random() * (cw - Inches(0.45))
            cy = yy + Inches(0.1) + rnd.random() * (hh - Inches(0.3))
            _box(sh, cx, cy, Inches(0.17), Inches(0.13), fill='FFFFFF', line=CELL, shape=MSO_SHAPE.OVAL,
                 name='Клетка')
        # деление: столбики плоских клеток
        yy, hh = ys[1]
        ncol = 7
        for c in range(ncol):
            cx = x0 + Inches(0.14) + c * (cw - Inches(0.28)) / ncol
            k, cy = 0, yy + Inches(0.06)
            while cy + Inches(0.1) < yy + hh - Inches(0.04):
                _box(sh, cx, cy, Inches(0.3), Inches(0.09), fill='FFFFFF', line=CELL, shape=MSO_SHAPE.OVAL,
                     name='Клетка')
                cy += Inches(0.115)
                k += 1
        # созревание: крупные клетки
        yy, hh = ys[2]
        for c in range(ncol):
            cx = x0 + Inches(0.12) + c * (cw - Inches(0.24)) / ncol
            cy = yy + Inches(0.06)
            while cy + Inches(0.3) < yy + hh:
                _box(sh, cx, cy, Inches(0.33), Inches(0.3), fill='FFFFFF', line=CELL, shape=MSO_SHAPE.OVAL,
                     name='Клетка')
                cy += Inches(0.34)
        # кость: перекладины
        yy, hh = ys[3]
        for c in range(ncol):
            cx = x0 + Inches(0.2) + c * (cw - Inches(0.4)) / ncol
            _line(sh, cx, yy + Inches(0.05), cx + Inches(0.12), yy + hh - Inches(0.05), col=BONE_LN, w=2.5,
                  name='Кость')
        _line(sh, x0, ys[3][0], x0 + cw, ys[3][0], col=C['axis'], w=1.0, name='Граница хряща и кости')
    # подписи зон и действие гормонов
    lx = x0 + cw + Inches(0.2)
    zx = lx + Inches(2.55)
    zw = MX + CONTENT_W - zx
    for i, (yy, hh) in enumerate(ys):
        z = zones[i]
        _line(sh, x0 + cw, yy + hh / 2, lx - Inches(0.03), yy + hh / 2, col=C['axis'], w=0.75)
        _text(sh, lx, yy, Inches(2.45), hh, z['name'], size=15, color=C['ink'], anchor=MSO_ANCHOR.MIDDLE,
              name='Зона')
        if z.get('acts'):
            _text(sh, zx, yy, zw, hh, z['acts'], size=15, anchor=MSO_ANCHOR.MIDDLE, name='Действие', gap=2)
        if i < 3:
            _line(sh, lx, yy + hh, MX + CONTENT_W, yy + hh, col=C['rule'], w=0.5)
    if g.get('bottom'):
        _text(sh, MX, bot + Inches(0.12), CONTENT_W, BODY_BOTTOM - bot - Inches(0.1), g['bottom'], size=16,
              color=C['text'], name='Вывод схемы')


# ---------- уровни: где ломается ----------

def levels(slide, d):
    """Слева колонка уровней со стрелками вниз, справа что происходит на каждом уровне.
    Справа может стоять изображение."""
    sh = slide.shapes
    g = d['diagram']
    rows = g['levels']
    img = g.get('image')
    has_img = img and os.path.exists(os.path.join(HERE, '..', img['path']))
    top = BODY_Y + Inches(0.05)
    bot = BODY_BOTTOM - (Inches(0.7) if g.get('bottom') else Inches(0.1))
    n = len(rows)
    gap = Inches(0.3)
    rh = (bot - top - gap * (n - 1)) / n
    lw = Inches(2.7)
    right_edge = MX + CONTENT_W - (Inches(3.6) if has_img else 0)
    tx = MX + lw + Inches(0.35)
    tw = right_edge - tx
    for i, r in enumerate(rows):
        y = top + i * (rh + gap)
        _box(sh, MX, y, lw, rh, fill=r.get('fill') or C['ink'], name=f'Уровень: {r["name"]}')
        tb = _text(sh, MX + Inches(0.15), y, lw - Inches(0.3), rh, r.get('sub') or [], size=13,
                   color=C['onink2'], anchor=MSO_ANCHOR.MIDDLE, head=r['name'], head_size=17, name='Уровень')
        for run in tb.text_frame.paragraphs[0].runs:
            run.font.color.rgb = _rgb(C['onink'])
        if i < n - 1:
            _arrow(sh, MX + lw / 2, y + rh, MX + lw / 2, y + rh + gap, col=C['axis'], w=2.0)
        _line(sh, MX + lw, y + rh / 2, tx - Inches(0.08), y + rh / 2, col=C['axis'], w=0.75)
        _text(sh, tx, y, tw, rh, r['text'], size=r.get('size', 16), anchor=MSO_ANCHOR.MIDDLE, name='Что происходит',
              gap=3)
    if has_img:
        _image(sh, img['path'], right_edge + Inches(0.3), top, Inches(3.3), bot - top, img.get('credit'))
    if g.get('bottom'):
        _text(sh, MX, bot + Inches(0.15), CONTENT_W, BODY_BOTTOM - bot - Inches(0.15), g['bottom'], size=16,
              name='Вывод схемы')


# ---------- эстрогены: путь и три сценария ----------

def estrogen(slide, d):
    sh = slide.shapes
    g = d['diagram']
    top = BODY_Y + Inches(0.05)
    # путь слева направо
    steps = g['path']                  # [{'name','text'}] 3–4 шага
    n = len(steps)
    gap = Inches(0.45)
    bw = (CONTENT_W - gap * (n - 1)) / n
    bh = Inches(1.35)
    for i, st in enumerate(steps):
        x = MX + i * (bw + gap)
        _box(sh, x, top, bw, bh, fill=C['ink'] if i == 0 else C['paper2'], line=None if i == 0 else C['rule'],
             name=f'Шаг: {st["name"]}')
        tb = _text(sh, x + Inches(0.12), top + Inches(0.08), bw - Inches(0.24), bh - Inches(0.16),
                   st.get('text') or [], size=13, color=C['onink2'] if i == 0 else C['text'], head=st['name'],
                   head_size=16, anchor=MSO_ANCHOR.MIDDLE, name='Шаг')
        if i == 0:
            for run in tb.text_frame.paragraphs[0].runs:
                run.font.color.rgb = _rgb(C['onink'])
        if i < n - 1:
            _arrow(sh, x + bw + Inches(0.05), top + bh / 2, x + bw + gap - Inches(0.05), top + bh / 2, w=2.0)
    # сценарии
    sc = g['scenarios']                # [{'name','text'}]
    y = top + bh + Inches(0.35)
    m = len(sc)
    sw = (CONTENT_W - Inches(0.3) * (m - 1)) / m
    sh_h = Inches(1.45)
    for i, it in enumerate(sc):
        x = MX + i * (sw + Inches(0.3))
        _line(sh, x, y, x + sw, y, col=C['height'], w=2.0, name='Сценарий')
        _text(sh, x, y + Inches(0.08), sw, sh_h, it['text'], size=14, head=it['name'], head_size=16,
              name='Сценарий', gap=2)
    if g.get('bottom'):
        by = y + sh_h + Inches(0.15)
        _text(sh, MX, by, CONTENT_W, BODY_BOTTOM - by, g['bottom'], size=15, name='Вывод схемы', gap=3)


# ---------- ген SHOX ----------

def shox(slide, d):
    sh = slide.shapes
    g = d['diagram']
    top = BODY_Y + Inches(0.1)
    img = g.get('image')
    has_img = img and os.path.exists(os.path.join(HERE, '..', img['path']))
    # две X-хромосомы и одна: слева
    def chrom(x, y, label, faded=False):
        w, h = Inches(0.42), Inches(2.6)
        p_h = h * 0.38
        col = C['centile'] if faded else C['height']
        _box(sh, x, y, w, p_h, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name='Короткое плечо')
        _box(sh, x, y + p_h + Inches(0.06), w, h - p_h - Inches(0.06), fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             name='Длинное плечо')
        if not faded:
            _box(sh, x - Inches(0.03), y + Inches(0.05), w + Inches(0.06), Inches(0.22), fill=C['weight'],
                 name='SHOX')
        _text(sh, x - Inches(0.3), y + h + Inches(0.08), w + Inches(0.6), Inches(0.3), label, size=13,
              color=C['muted'], align=PP_ALIGN.CENTER, name='Подпись хромосомы')
    cx = MX + Inches(0.1)
    chrom(cx, top + Inches(0.35), 'X')
    chrom(cx + Inches(0.7), top + Inches(0.35), 'X')
    _text(sh, cx - Inches(0.1), top - Inches(0.1), Inches(1.6), Inches(0.35), g['left_label'], size=14,
          color=C['ink'], align=PP_ALIGN.CENTER, name='Норма')
    cx2 = MX + Inches(2.1)
    chrom(cx2, top + Inches(0.35), 'X')
    chrom(cx2 + Inches(0.7), top + Inches(0.35), '', faded=True)
    _text(sh, cx2 - Inches(0.1), top - Inches(0.1), Inches(1.6), Inches(0.35), g['right_label'], size=14,
          color=C['ink'], align=PP_ALIGN.CENTER, name='Тернер')
    # подпись гена
    _line(sh, cx2 + Inches(0.42), top + Inches(0.5), cx2 + Inches(1.6), top + Inches(0.5), col=C['weight'], w=1.0)
    _text(sh, cx2 + Inches(1.65), top + Inches(0.3), Inches(1.8), Inches(0.45), g['gene_label'], size=13,
          color=C['ink'], name='Ген')
    # текст справа
    tx = MX + Inches(4.1)
    tw = MX + CONTENT_W - tx - (Inches(3.5) if has_img else 0)
    _text(sh, tx, top + Inches(0.6), tw, BODY_BOTTOM - top - Inches(0.7), g['points'], size=16, name='Механизм',
          gap=8)
    if has_img:
        _image(sh, img['path'], MX + CONTENT_W - Inches(3.3), top, Inches(3.3), BODY_BOTTOM - top - Inches(0.1),
               img.get('credit'))


DRAW = {'growth_plate': growth_plate, 'levels': levels, 'estrogen': estrogen, 'shox': shox}
