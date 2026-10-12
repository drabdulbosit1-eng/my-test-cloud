"""Схематичные кривые роста как редактируемые фигуры PowerPoint.

Центильные кривые построены по модели Preece–Baines 1 с условным СО.
Это схема для слайда, а не референсные данные.
"""
import math
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from style import C, FONT_BODY

# Preece–Baines 1: h1, h_theta, s0, s1, theta
PB = {
    'M': (174.5, 162.9, 0.115, 1.20, 14.4),
    'F': (162.4, 151.8, 0.110, 1.20, 12.0),
}
# медиана массы, кг (схема)
W50 = {
    'M': [(2, 12.5), (5, 18.5), (8, 25.5), (10, 32), (12, 40), (14, 51), (16, 61), (18, 68)],
    'F': [(2, 12.0), (5, 18.0), (8, 25.0), (10, 32.5), (12, 41), (14, 50), (16, 54), (18, 56)],
}


def h50(sex, t):
    h1, ht, s0, s1, th = PB[sex]
    return h1 - 2 * (h1 - ht) / (math.exp(s0 * (t - th)) + math.exp(s1 * (t - th)))


def hsd(t):
    return min(6.5, 3.6 + 0.2 * max(0.0, t - 2))


def height(sex, t, z):
    return h50(sex, t) + z * hsd(t)


def _interp(pts, t):
    for (a, va), (b, vb) in zip(pts, pts[1:]):
        if a <= t <= b:
            return va + (vb - va) * (t - a) / (b - a)
    return pts[0][1] if t < pts[0][0] else pts[-1][1]


def weight(sex, t, z):
    m = _interp(W50[sex], t)
    # асимметрия: вверх шире, чем вниз
    k = 0.11 if z < 0 else 0.17
    return m * math.exp(k * z)


def _rgb(hexs):
    return RGBColor.from_string(hexs)


class Panel:
    """Прямоугольная область графика: перевод (возраст, значение) в EMU."""

    def __init__(self, x, y, w, h, t0, t1, v0, v1):
        self.x, self.y, self.w, self.h = x, y, w, h
        self.t0, self.t1, self.v0, self.v1 = t0, t1, v0, v1

    def px(self, t):
        return int(self.x + (t - self.t0) / (self.t1 - self.t0) * self.w)

    def py(self, v):
        return int(self.y + self.h - (v - self.v0) / (self.v1 - self.v0) * self.h)


def _line(shapes, pts, color, width_pt, dash=None, name=None):
    pts = [(int(a), int(b)) for a, b in pts]
    fb = shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    fb.add_line_segments(pts[1:], close=False)
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = _rgb(color)
    shp.line.width = Pt(width_pt)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shp.line.dash_style = dash
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def _poly(shapes, pts, fill, name=None):
    pts = [(int(a), int(b)) for a, b in pts]
    fb = shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    fb.add_line_segments(pts[1:], close=True)
    shp = fb.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def _label(shapes, x, y, w, h, text, size=11, color=None, bold=False, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.MIDDLE, italic=False):
    tb = shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.name = FONT_BODY
    f.bold = bold
    f.italic = italic
    f.color.rgb = _rgb(color or C['muted'])
    return tb


def _steps(t0, t1, n=60):
    return [t0 + (t1 - t0) * i / n for i in range(n + 1)]


def centile_panel(shapes, x, y, w, h, sex, t0, t1, kind, patient, events=(), label=None,
                  pcolor=None, axis_labels=True):
    """kind: 'height' или 'weight'. patient: [(возраст, z)] — траектория пациента в СО."""
    fn = height if kind == 'height' else weight
    ts = _steps(t0, t1)
    lo = [fn(sex, t, -1.88) for t in ts]
    hi = [fn(sex, t, 1.88) for t in ts]
    pz = [fn(sex, a, z) for a, z in patient]
    vmin = min(lo + pz)
    vmax = max(hi + pz)
    pad = (vmax - vmin) * 0.06
    P = Panel(x, y, w, h, t0, t1, vmin - pad, vmax + pad)

    # полоса 3–97-го центиля
    band = [(P.px(t), P.py(v)) for t, v in zip(ts, hi)] + \
           [(P.px(t), P.py(v)) for t, v in reversed(list(zip(ts, lo)))]
    _poly(shapes, band, C['band'], name=f'{kind}: полоса 3–97')
    for z, nm in ((-1.88, '3'), (0, '50'), (1.88, '97')):
        pts = [(P.px(t), P.py(fn(sex, t, z))) for t in ts]
        _line(shapes, pts, C['centile'], 0.75, name=f'{kind}: {nm}-й центиль')
        if axis_labels:
            _label(shapes, P.px(t1) + Emu(40000), P.py(fn(sex, t1, z)) - Emu(80000), Emu(300000),
                   Emu(160000), nm, size=9)

    # ось возраста
    _line(shapes, [(P.x, P.y + P.h), (P.x + P.w, P.y + P.h)], C['axis'], 0.75, name=f'{kind}: ось')
    if axis_labels:
        step = 2 if (t1 - t0) > 8 else 1
        a = math.ceil(t0)
        while a <= t1:
            if (a - math.ceil(t0)) % step == 0:
                _label(shapes, P.px(a) - Emu(150000), P.y + P.h + Emu(30000), Emu(300000), Emu(170000),
                       str(int(a)), size=9, align=PP_ALIGN.CENTER)
            a += 1

    # траектория пациента
    pts = []
    for (a0, z0), (a1, z1) in zip(patient, patient[1:]):
        for t in _steps(a0, a1, 12)[:-1]:
            z = z0 + (z1 - z0) * (t - a0) / (a1 - a0)
            pts.append((P.px(t), P.py(fn(sex, t, z))))
    a, z = patient[-1]
    pts.append((P.px(a), P.py(fn(sex, a, z))))
    col = pcolor or (C['height'] if kind == 'height' else C['weight'])
    _line(shapes, pts, col, 2.75, name=f'{kind}: пациент')
    for a, z in patient:
        d = Emu(60000)
        dot = shapes.add_shape(9, P.px(a) - d, P.py(fn(sex, a, z)) - d, 2 * d, 2 * d)  # oval
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(col)
        dot.line.fill.background()
        dot.shadow.inherit = False
        dot.name = f'{kind}: визит {a}'

    # события: вертикальная пунктирная метка
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    for i, ev in enumerate(events):
        a, text = ev[0], ev[1]
        _line(shapes, [(P.px(a), P.y), (P.px(a), P.y + P.h)], C['muted'], 0.75,
              dash=MSO_LINE_DASH_STYLE.DASH, name=f'событие {a}')
        _label(shapes, P.px(a) + Emu(50000), P.y + i * Emu(190000), Emu(1600000), Emu(190000), text, size=10,
               color=C['text'], anchor=MSO_ANCHOR.TOP)

    if label:
        _label(shapes, P.x, P.y - Emu(230000), P.w, Emu(200000), label, size=12, color=C['text'],
               bold=True, anchor=MSO_ANCHOR.BOTTOM)
    return P


def growth_chart(shapes, x, y, w, h, sex, t0, t1, height_z, weight_z, events=(), title=None):
    """Два стыкованных поля: рост (сверху, 62 %) и масса (снизу)."""
    gap = Emu(330000)
    top_h = int((h - gap) * 0.62)
    bot_h = int(h - gap - top_h)
    off = Emu(240000) if title else 0
    if title:
        _label(shapes, x, y, w, Emu(220000), title, size=13, color=C['ink'], bold=True,
               anchor=MSO_ANCHOR.TOP)
    centile_panel(shapes, x, y + off + Emu(230000), w - Emu(320000), top_h - off - Emu(230000), sex, t0, t1,
                  'height', height_z, events, label='Рост, см')
    centile_panel(shapes, x, y + top_h + gap + Emu(230000), w - Emu(320000), bot_h - Emu(230000), sex, t0,
                  t1, 'weight', weight_z, (), label='Масса, кг')


def z_thumb(shapes, x, y, w, h, title, t0, t1, hz, wz, note=None):
    """Миниатюра: рост и масса в СО на одной оси (−3…+3)."""
    _label(shapes, x, y, w, Emu(260000), title, size=13, color=C['ink'], bold=True, anchor=MSO_ANCHOR.TOP)
    py0 = y + Emu(300000)
    ph = h - Emu(300000) - (Emu(480000) if note else Emu(220000))
    P = Panel(x + Emu(260000), py0, w - Emu(300000), ph, t0, t1, -3.2, 3.2)
    band = [(P.px(t0), P.py(1.88)), (P.px(t1), P.py(1.88)), (P.px(t1), P.py(-1.88)), (P.px(t0), P.py(-1.88))]
    _poly(shapes, band, C['band'], name=f'{title}: норма')
    _line(shapes, [(P.px(t0), P.py(0)), (P.px(t1), P.py(0))], C['centile'], 0.75, name=f'{title}: 0')
    for v in (-2, 0, 2):
        _label(shapes, x, P.py(v) - Emu(80000), Emu(230000), Emu(160000),
               ('+' if v > 0 else '') + str(v).replace('-', '−'), size=9, align=PP_ALIGN.RIGHT)
    for series, col, nm in ((wz, C['weight'], 'масса'), (hz, C['height'], 'рост')):
        pts = []
        for (a0, z0), (a1, z1) in zip(series, series[1:]):
            for t in _steps(a0, a1, 10)[:-1]:
                pts.append((P.px(t), P.py(z0 + (z1 - z0) * (t - a0) / (a1 - a0))))
        pts.append((P.px(series[-1][0]), P.py(series[-1][1])))
        _line(shapes, pts, col, 2.5, name=f'{title}: {nm}')
    _label(shapes, P.x, P.y + P.h + Emu(20000), Emu(500000), Emu(160000), f'{t0:g} лет', size=9)
    _label(shapes, P.x + P.w - Emu(500000), P.y + P.h + Emu(20000), Emu(500000), Emu(160000), f'{t1:g} лет',
           size=9, align=PP_ALIGN.RIGHT)
    if note:
        _label(shapes, x, y + h - Emu(250000), w, Emu(250000), note, size=11, color=C['text'],
               anchor=MSO_ANCHOR.TOP)
