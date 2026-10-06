"""Promo rework, slides 17-24 of base47.pptx (old numbering), source: promo/final_17_24.md.
Build: cd W && PYTHONPATH=W:W/promo python3 apply14.py base47.pptx promo/t_17.pptx impl_17_24
One function per slide: s17(prs) ... s24(prs); run(prs) calls them in order."""
import copy
import io

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import measure as M
import r14lib as L
from lib import (add_box, add_text, bar, card, chevron, delete_shape, get_rpr, iter_shapes,
                 style_rpr)

BLACK = "000000"
PURPLE = "7030A0"
GREY = "7F7F7F"      # placebo
CYAN = "0F9ED5"      # comparator / SGLT2 inhibitor
GREEN = "4EA72E"     # combination
TEAL = "156082"      # neutral plate
ORANGE = "E97132"    # potassium plate
CARD = "F2F2F2"
CHEV = "D9D9D9"
BODY_BOTTOM = 538    # bottom edge of the footer text in the source deck


# ----------------------------------------------------------------------------- helpers
def geom(sh, x=None, y=None, w=None, h=None):
    if x is not None:
        sh.left = Pt(x)
    if y is not None:
        sh.top = Pt(y)
    if w is not None:
        sh.width = Pt(w)
    if h is not None:
        sh.height = Pt(h)
    return sh


def get_geom(sh):
    return Emu(sh.left).pt, Emu(sh.top).pt, Emu(sh.width).pt, Emu(sh.height).pt


def by_name(slide, name):
    """All shapes (groups included) with exactly this name, in z-order."""
    return [sh for sh in iter_shapes(slide.shapes) if sh.name == name]


def by_prefix(slide, prefix):
    return [sh for sh in iter_shapes(slide.shapes) if sh.name.startswith(prefix)]


def one(slide, name):
    r = by_name(slide, name)
    if len(r) != 1:
        raise KeyError(f"{name!r}: {len(r)} shapes")
    return r[0]


def recolor_text(sh, color, bold=None, size=None):
    for r in sh._element.iter(qn("a:r")):
        style_rpr(get_rpr(r), color=color, bold=bold, size=size)


def recolor_fill(sh, color):
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)


def recolor_line(sh, color, width=None):
    sh.line.color.rgb = RGBColor.from_string(color)
    if width is not None:
        sh.line.width = Pt(width)


def shift(shapes, dx=0.0, dy=0.0):
    for sh in shapes:
        sh.left = Emu(int(sh.left + dx * 12700))
        sh.top = Emu(int(sh.top + dy * 12700))


def set_line_spacing(sh, pct):
    for p in sh._element.iter(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        if ppr is None:
            ppr = etree.SubElement(p, qn("a:pPr"))
            p.insert(0, ppr)
        for k in ppr.findall(qn("a:lnSpc")):
            ppr.remove(k)
        ln = etree.Element(qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct")).set("val", str(int(pct * 1000)))
        ppr.insert(0, ln)


def rich(sh, paras, size=None, color=BLACK):
    """Rewrite a text shape: paras = list of paragraphs, each a list of (text, bold) runs.
    Keeps the first run format of the shape (font, size); `size` overrides it, `color` recolours."""
    L.set_rich_lines(sh, paras, color=color)
    if size is not None:
        for r in sh._element.iter(qn("a:r")):
            style_rpr(get_rpr(r), size=size)


def lines_of(paras, w, size):
    """Number of lines of paragraphs [[(text, bold), ...], ...] in width w at `size` pt."""
    return sum(M.n_lines(p, w, size) for p in paras)


def txt(slide, x, y, w, h, paras, name, anchor="t", spacing=0.9, after=None):
    """Text box; paras = list of paragraphs, each a list of runs (text, size, bold, color)."""
    return add_text(slide, x, y, w, h, paras, anchor=anchor, name=name, line_spacing=spacing,
                    space_after=after)


def plate(slide, x, y, w, h, color, text, name, size=17):
    """Coloured plate with white text; '\n' in `text` starts a new paragraph (second line)."""
    lines = text.split("\n")
    sh = bar(slide, x, y, w, h, color, lines[0], size=size, name=name)
    for extra in lines[1:]:
        p0 = sh.text_frame.paragraphs[0]._p
        p = copy.deepcopy(p0)
        p0.getparent().append(p)
        p.find(qn("a:r")).find(qn("a:t")).text = extra
    return sh


def grey_card(slide, x, y, w, h, name):
    return card(slide, x, y, w, h, name=name)


def footers(slide, abbr, sources, abbr_size=10, abbr_lines=None, abbr_w=850):
    """Rewrite 'Источники' (8 pt, bottom edge fixed, grows upwards) and 'Сокращения' (10 pt, sits
    above the sources). Both black. Returns the y above which body content has to end."""
    src = L.shape_named(slide, "Источники")
    L.set_lines(src, sources)
    n = sum(M.n_lines([(p, False)], 850, 8) for p in sources)
    h = max(16.0, round(n * 9.8 + 2, 1))
    geom(src, y=BODY_BOTTOM - h, h=h)
    ab = L.shape_named(slide, "Сокращения")
    L.set_lines(ab, [abbr])
    na = abbr_lines or M.n_lines([(abbr, False)], abbr_w, abbr_size)
    ha = round(na * 12.4, 1)
    geom(ab, w=abbr_w)
    top = min(494.0, BODY_BOTTOM - h - 3 - ha)
    geom(ab, y=top, h=ha)
    for sh in (src, ab):
        for r in sh._element.iter(qn("a:r")):
            style_rpr(get_rpr(r), color=BLACK)
    return top


def flow(slide, x, y, w, paras, name, size=14, spacing=0.9, after=3, color=BLACK):
    """One text box with paragraphs [[(text, bold), ...], ...]; returns (shape, text height)."""
    plain = [[(r[0], r[1]) for r in p] for p in paras]
    n = lines_of(plain, w, size)
    h = n * size * 1.15 * spacing + after * (len(paras) - 1)
    runs = [[(r[0], size, r[1], r[2] if len(r) > 2 else color) for r in p] for p in paras]
    tb = txt(slide, x, y, w, h + 2, runs, name, spacing=spacing, after=after)
    return tb, h


def chart_transform(slide, prefix, old_o, new_o, kx=1.0, ky=1.0):
    """Move/scale a hand-drawn chart (lines, ticks, labels, curves) about its axis origin.
    old_o / new_o = (x of the y axis, y of the x axis). Text sizes stay, positions are mapped."""
    import re
    ox, oy = old_o
    nx, ny = new_o

    def X(x):
        return nx + (x - ox) * kx

    def Y(y):
        return ny + (y - oy) * ky

    for sh in list(iter_shapes(slide.shapes)):
        n = sh.name
        if n.startswith("published_curve_" + prefix) and prefix:
            geom(sh, x=nx, y=Y(sh.top / 12700), w=Emu(sh.width).pt * kx, h=Emu(sh.height).pt * ky)
            continue
        if not n.startswith(prefix):
            continue
        k = n[len(prefix):]
        x, y, w, h = get_geom(sh)
        if k in ("y_axis_title", "x_axis_title") or k.startswith("legend"):
            continue
        if k == "y_axis":
            geom(sh, x=nx, y=Y(y), h=h * ky)
        elif k == "x_axis":
            geom(sh, x=nx, y=ny, w=w * kx)
        elif k.startswith("tick_y"):
            geom(sh, x=nx - w, y=Y(y))
        elif k.startswith("tick_x"):
            geom(sh, x=X(x), y=ny + (y - oy))
        elif k.startswith("grid_"):
            geom(sh, x=nx, y=Y(y), w=w * kx)
        elif re.match(r"y_\d+$", k):
            geom(sh, x=nx - 6 - w, y=Y(y + h / 2) - h / 2)
        elif k.startswith("month_"):
            geom(sh, x=X(x + w / 2) - w / 2, y=ny + (y - oy))
        elif k.startswith("published_curve"):
            geom(sh, x=nx, y=Y(y), w=w * kx, h=h * ky)


def numbers_row(slide, x, y, w, fin, pla, name):
    """Two big numbers (28 pt) with the group names: finerenone purple, placebo grey."""
    a = txt(slide, x, y, w / 2, 34, [[(fin, 28, True, PURPLE), ("  Финеренон", 14, False, PURPLE)]],
            f"{name}_финеренон", spacing=1.0)
    b = txt(slide, x + w / 2, y, w / 2, 34, [[(pla, 28, True, GREY), ("  Плацебо", 14, False, GREY)]],
            f"{name}_плацебо", spacing=1.0)
    return a, b


def style_curves(slide, prefix=""):
    """Chart colours of the course: finerenone #7030A0 solid, placebo #7F7F7F dashed."""
    for nm in ("published_curve_" + prefix + "Финеренон", prefix + "legend_finerenone"):
        recolor_line(one(slide, nm), PURPLE)
    for nm in ("published_curve_" + prefix + "Плацебо", prefix + "legend_placebo"):
        recolor_line(one(slide, nm), GREY)
    recolor_text(one(slide, prefix + "legend_finerenone_label"), PURPLE)
    recolor_text(one(slide, prefix + "legend_placebo_label"), GREY)


def notes(slide, *paras):
    """Speaker notes: paragraphs separated by an empty paragraph (as in the source deck)."""
    L.set_notes(slide, "\n\n".join(paras))


def copy_pic(src_shape, dst_slide, x, y, w, h, name):
    pic = dst_slide.shapes.add_picture(io.BytesIO(src_shape.image.blob), Pt(x), Pt(y), Pt(w), Pt(h))
    pic.name = name
    return pic


def check_fit(label, paras, w, size, h, spacing=0.9):
    """Warn when text does not fit (Arial Narrow metrics); line height = size * 1.15 * spacing."""
    n = lines_of(paras, w, size)
    need = n * size * 1.15 * spacing
    flag = "OK " if need <= h + 0.5 else "!! "
    print(f"  fit {flag}{label}: {n} lines, {need:.0f} pt of {h:.0f} pt")
    return need


# ----------------------------------------------------------------------------- slide 17
def s17(prs):
    s = L.S(prs, 17)
    # title
    L.set_title(s, "Избыточная активация МКР поддерживает воспаление и фиброз в почке и сердце")

    # scheme goes up by 27 pt (gap between the title and the scheme); everything between y 100 and 396
    dy = -27
    moving = [sh for sh in iter_shapes(s.shapes) if 100 <= Emu(sh.top).pt < 396
              and sh.name not in ("Информативный заголовок", "Логотип Орвилле")]
    shift(moving, dy=dy)

    # colours: protein of the receptor grey (as on the next slide), labels black, aldosterone blue
    for sh in by_name(s, "Контур условного символа рецептора"):
        recolor_fill(sh, GREY)
    recolor_fill(one(s, "Блик условного символа рецептора"), "B3B3B3")
    recolor_text(one(s, "Ядро"), BLACK)
    recolor_text(one(s, "Рецептор"), BLACK)
    recolor_text(one(s, "Усиливаются сигналы воспаления и фиброза"), BLACK, bold=True)
    recolor_text(one(s, "Альдостерон"), TEAL, bold=True)
    recolor_text(one(s, "Метка_фибробласты"), BLACK)
    recolor_line(one(s, "Подпись к альдостерону"), TEAL)
    recolor_line(one(s, "Активация рецептора связана с регуляцией генов"), BLACK)

    # red arrow cell -> tissue becomes a wide grey chevron with the label "Тканевой ответ"
    delete_shape(one(s, "Связь клеточного сигнала с тканевым ответом"))
    lab = one(s, "Тканевой ответ")
    geom(lab, x=440, y=181, w=68, h=34)
    chevron(s, 462, 222, w=30, h=46, name="Шеврон_тканевой_ответ", color=CHEV)

    # captions under the scheme: text, cards for inflammation / fibrosis
    for nm in ("Воспаление", "Приток и активация иммунных клеток.", "Фиброз",
               "Накопление матрикса между клетками.", "Связь_с_финереноном"):
        delete_shape(one(s, nm))
    base = L.shape_with(s, "При диабете и повреждении")
    base.name = "Пояснение_МКР"
    rich(base, [[("При диабете и ХБП сигнал МКР может усиливаться: рецептор сильнее включает "
                  "гены воспаления и фиброза.", False)]], size=16)
    geom(base, x=40, y=412, w=440, h=40)
    y_bar, h_bar, y_card, h_card = 368, 28, 398, 72
    for x, title, body, nm in (
            (508, "Воспаление", "Приток и активация иммунных клеток.", "воспаление"),
            (721, "Фиброз", "Рубцовая ткань: избыток межклеточного вещества (матрикса).", "фиброз")):
        plate(s, x, y_bar, 200, h_bar, TEAL, title, f"Полоса_{nm}")
        grey_card(s, x, y_card, 200, h_card, f"Карточка_{nm}")
        txt(s, x + 10, y_card + 6, 180, h_card - 12, [[(body, 16, False, BLACK)]], f"Текст_{nm}")
        check_fit(nm, [[(body, False)]], 180, 16, h_card - 12)

    # small organ icons: kidney and heart (same pictures as on the study-programme slide)
    s19 = L.S(prs, 19)
    kid = [sh for sh in s19.shapes if sh.name == "Рисунок 44"][0]
    heart = [sh for sh in s19.shapes if sh.name == "Рисунок 45"][0]
    copy_pic(kid, s, 856, 94, 30, 46, "Иконка_почка")
    copy_pic(heart, s, 890, 94, 34, 46, "Иконка_сердце")
    lab2 = txt(s, 700, 106, 148, 22, [[("Почка и сердце", 14, False, BLACK)]], "Подпись_почка_и_сердце")
    lab2.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    footers(s,
            "МКР: минералокортикоидный рецептор; ХБП: хроническая болезнь почек.",
            ["Обзоры: Agarwal R, et al. Eur Heart J. 2021;42:152–161. Barrera-Chimal J, et al. Kidney Int. "
             "2019;96:302–319. Иллюстрации: Servier Medical Art (CC BY 4.0)."])

    notes(s,
          "На прошлом слайде рецептор работал в норме: регулировал натрий и калий в клетке канальца. "
          "Теперь о том, что происходит, когда он активен сильнее нормы.",
          "При диабете и хронической болезни почек сигнал минералокортикоидного рецептора может усиливаться. "
          "Такая избыточная активация способна поддерживать воспаление и накопление межклеточного матрикса. "
          "На схеме показаны разные клетки: иммунные, эпителиальные и фибробласты. Фибробласты образуют "
          "компоненты матрикса. Его избыточное накопление связано с фиброзом. На схеме фрагмент почки. "
          "Рецептор есть и в сердце, и в сосудах, и там его избыточная активация тоже поддерживает "
          "воспаление и фиброз.",
          "Если этот сигнал заблокировать, воспаление и фиброз могут стать слабее. Как это делает "
          "финеренон, показывает следующий слайд.",
          "Фраза для визита: «При диабете и ХБП минералокортикоидный рецептор может быть избыточно "
          "активен и поддерживать воспаление и фиброз в почке и сердце. Ингибиторы АПФ, БРА и иНГЛТ2 сам "
          "этот рецептор не блокируют».",
          "Если врач спросит: «Доказана ли эта связь у людей?» Связь избыточной активации МКР с воспалением "
          "и фиброзом описана в экспериментальных работах и обзорах (Agarwal, 2021). Эффект финеренона на "
          "воспаление и фиброз показан в доклинических моделях. В клинических исследованиях оценивали "
          "исходы, они показаны дальше.")


# ----------------------------------------------------------------------------- slide 18
def s18(prs):
    s = L.S(prs, 18)
    L.set_title(s, "Финеренон избирательно блокирует МКР, и гены воспаления и фиброза включаются слабее")

    # scheme and columns go up by 27 pt
    shift([sh for sh in iter_shapes(s.shapes) if 100 <= Emu(sh.top).pt < 400
           and sh.name not in ("Информативный заголовок", "Логотип Орвилле")], dy=-27)

    recolor_text(one(s, "Ядро"), BLACK)
    inside = L.shape_with(s, "Сигналы воспаления и фиброза")
    L.set_lines(inside, ["Слабее включаются гены", "воспаления и фиброза"])
    geom(inside, y=318, h=40)

    # class of the drug (text under the cell)
    cls = L.shape_with(s, "Класс: нестероидный")
    rich(cls, [[("Нестероидный избирательный антагонист МКР. Связывается с МКР внутри клетки и блокирует "
                 "привлечение белков-коактиваторов, которые нужны для включения генов воспаления и "
                 "фиброза.", False)]])
    geom(cls, x=40, y=406, w=440, h=72)
    check_fit("класс", [[("Нестероидный избирательный антагонист МКР. Связывается с МКР внутри клетки и блокирует "
                          "привлечение белков-коактиваторов, которые нужны для включения генов воспаления и "
                          "фиброза.", False)]], 440, 17, 72)

    # right column: two plates with cards
    L.rep(s, "Клиническая польза", "Доказано при ХБП с альбуминурией и диабете 2 типа")
    L.rep(s, "Почему контролируют калий", "Калий: правило простое")
    y1 = 97
    geom(one(s, "Полоса_польза"), y=y1, h=32)
    geom(one(s, "Карточка_польза"), y=y1 + 36, h=136)
    benefit = L.shape_with(s, "Доказанная польза")
    rich(benefit, [[("Финеренон на фоне ИАПФ или БРА в максимально переносимой дозе снижал по сравнению с "
                     "плацебо риск почечного исхода (FIDELIO-DKD) и сердечно-сосудистого исхода "
                     "(FIGARO-DKD).", False)]])
    geom(benefit, y=y1 + 44, h=76)
    studies = L.shape_with(s, "Исследования:")
    rich(studies, [[("Исследования: ", True), ("FIDELIO-DKD, FIGARO-DKD; объединённый анализ FIDELITY", False)]])
    geom(studies, y=y1 + 36 + 136 - 8 - 32, h=32)
    y2 = y1 + 36 + 136 + 12            # 281
    geom(one(s, "Полоса_калий"), y=y2, h=32)
    geom(one(s, "Карточка_калий"), y=y2 + 36, h=134)
    potassium = L.shape_with(s, "Блокада МКР может уменьшать")
    rich(potassium, [[("Блокада МКР в почечных канальцах снижает выведение калия, поэтому его уровень в "
                       "крови может расти.", False)],
                     [("Калий и рСКФ проверяют до начала лечения, через 4 недели после начала или "
                       "повышения дозы и далее периодически.", False)]])
    geom(potassium, y=y2 + 44, h=118)
    for p in potassium.text_frame.paragraphs[1:]:
        p.space_before = Pt(6)

    footers(s,
            "МКР: минералокортикоидный рецептор; ХБП: хроническая болезнь почек; рСКФ: расчётная скорость "
            "клубочковой фильтрации; ИАПФ: ингибитор ангиотензинпревращающего фермента; БРА: блокатор "
            "рецепторов ангиотензина II.",
            ["Финеренон-Орвилле. Инструкция по медицинскому применению. Bakris GL, et al. N Engl J Med. "
             "2020;383:2219–2229. Pitt B, et al. N Engl J Med. 2021;385:2252–2263. Agarwal R, et al. Eur "
             "Heart J. 2022;43:474–484. Механизм и доклинические данные: Amazit L, et al. J Biol Chem. "
             "2015;290:21876–21889; Grune J, et al. Hypertension. 2018;71:599–608; Droebner K, et al. Am J "
             "Nephrol. 2021;52:588–601. Иллюстрации: Servier Medical Art (CC BY 4.0)."])

    notes(s,
          "На прошлом слайде рецептор работал избыточно. Теперь о том, что меняет финеренон.",
          "Финеренон связывается с минералокортикоидным рецептором внутри клетки и препятствует его "
          "активации. Поэтому его называют антагонистом рецептора. Избирательность означает преимущественное "
          "действие на эту мишень, а слово «нестероидный» описывает химическое строение. По инструкции "
          "финеренон блокирует привлечение коактиваторов: это белки-помощники, которые рецептору нужны, чтобы "
          "включить гены воспаления и фиброза. В экспериментах на животных финеренон уменьшал фиброз в "
          "сердце и почках.",
          "Клинические эффекты изучали у взрослых с хронической болезнью почек и диабетом 2 типа в "
          "FIDELIO-DKD, FIGARO-DKD и объединённом анализе FIDELITY. Финеренон сравнивали с плацебо на фоне "
          "ингибитора АПФ или блокатора рецепторов ангиотензина в максимально переносимой дозе. Блокада "
          "рецептора в почечных канальцах снижает выведение калия, поэтому его уровень в крови может "
          "повышаться. Правило простое: калий и рСКФ проверяют до начала лечения, через 4 недели после "
          "начала или повышения дозы и далее периодически. Помните схему с натрием и калием: когда рецептор "
          "заблокирован, калий хуже выходит в просвет канальца.",
          "Как блокада рецептора приводит к исходам, показывает следующий слайд: цепочка от рецептора до "
          "строк показания.",
          "Фраза для визита: «Финеренон блокирует минералокортикоидный рецептор, который при диабете и ХБП "
          "может работать избыточно. У взрослых с ХБП, альбуминурией и диабетом 2 типа на фоне ИАПФ или БРА "
          "он снизил по сравнению с плацебо риск почечных исходов (FIDELIO-DKD) и сердечно-сосудистых "
          "исходов (FIGARO-DKD). Калий проверяют до начала лечения, через 4 недели и далее периодически».",
          "Если врач спросит: «Влияние на воспаление и фиброз доказано у людей?» Данные получены в "
          "экспериментах: у мышей финеренон уменьшал фиброз сердца и инвазию макрофагов в ткань (Grune, "
          "2018); в моделях фиброза почки он уменьшал накопление миофибробластов и коллагена, а маркеры "
          "воспаления в тестированных дозах не менялись (Droebner, 2021). У пациентов оценивали клинические "
          "исходы, они на следующих слайдах.")


# ----------------------------------------------------------------------------- slide 19
def s19(prs):
    s = L.S(prs, 19)
    L.set_title(s, "Более 13 000 пациентов с ХБП и диабетом 2 типа: FIDELIO-DKD оценивал почечные исходы, "
                   "FIGARO-DKD сердечно-сосудистые")

    # left text -> two Levin blocks (plate #156082 + grey card)
    delete_shape(one(s, "Text 2"))
    blocks = [("Кто участвовал", "Пациенты с хронической болезнью почек, сахарным диабетом 2 типа и "
                                 "альбуминурией; калий сыворотки при включении ≤4,8 ммоль/л.", 90, "участники"),
              ("Что сравнивали", "Финеренон или плацебо на фоне ИАПФ или БРА в максимально переносимой "
                                 "дозе.", 76, "сравнение")]
    y = 90
    for title, body, h_card, nm in blocks:
        plate(s, 36, y, 276, 30, TEAL, title, f"Полоса_{nm}")
        grey_card(s, 36, y + 34, 276, h_card, f"Карточка_{nm}")
        txt(s, 46, y + 34 + 6, 256, h_card - 12, [[(body, 16, False, BLACK)]], f"Текст_{nm}")
        check_fit(nm, [[(body, False)]], 256, 16, h_card - 12)
        y += 34 + h_card + 14

    # colours: study names and primary outcome labels black bold, grey rules
    for nm in ("fidelio_title", "figaro_title", "fidelio_primary_label", "figaro_primary_label"):
        recolor_text(one(s, nm), BLACK, bold=True)
    for nm in ("trial_blue_rule", "trial_green_rule"):
        recolor_line(one(s, nm), CHEV, width=1.75)

    L.rep(s, "Ключевой вторичный исход:", "Вторичный исход:", count=2)
    L.rep(s, "Заранее запланированный поисковый анализ объединённых индивидуальных данных.",
          "Заранее запланированный объединённый анализ данных каждого пациента.")
    delete_shape(one(s, "safety_outcome"))
    # FIDELITY rows go up by 8 pt (the line about hyperkalaemia is gone)
    for nm in ("fidelity_title", "fidelity_n", "fidelity_description"):
        shift([one(s, nm)], dy=-8)
    center = one(s, "pooled_center")
    geom(center, h=3)

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; ИАПФ: ингибитор "
            "ангиотензинпревращающего фермента; БРА: блокатор рецепторов ангиотензина II.",
            [L.shape_named(s, "Источники").text_frame.text])

    notes(s,
          "Теперь посмотрим, что показали исследования у пациентов.",
          "В FIDELIO-DKD главным был комбинированный почечный исход. В FIGARO-DKD главным был "
          "сердечно-сосудистый исход. В обоих исследованиях участвовали взрослые с хронической болезнью "
          "почек и диабетом 2 типа. Критерии включения предусматривали альбуминурию и калий не выше 4,8 "
          "ммоль/л. Финеренон сравнивали с плацебо на фоне ингибитора АПФ или блокатора рецепторов "
          "ангиотензина в максимально переносимой дозе.",
          "FIDELITY объединил индивидуальные данные этих работ. Из 13 171 случайно распределённого участника "
          "в анализ вошли 13 026. Разница в 145 участников: их исключили до анализа из-за критических "
          "нарушений правил надлежащей клинической практики в исследовательских центрах.",
          "Объединённый анализ запланировали заранее; значения p в нём приведены без поправки на "
          "множественные сравнения. В его почечном исходе учитывали устойчивое снижение рСКФ не менее чем на "
          "57%, а в исходных исследованиях порог был 40%, поэтому почечные результаты FIDELITY и отдельных "
          "исследований напрямую не сравнивают.",
          "Начнём с FIDELIO-DKD: исследования, где главным исходом были почки.",
          "Фраза для визита: «Финеренон изучали в двух исследованиях III фазы, более чем у 13 000 пациентов "
          "с ХБП и диабетом 2 типа на фоне ИАПФ или БРА. FIDELIO-DKD оценивал почечные исходы, FIGARO-DKD "
          "сердечно-сосудистые, а FIDELITY объединил данные обоих».",
          "Если врач спросит: «А пациенты на иНГЛТ2?» В FIDELIO-DKD и FIGARO-DKD иНГЛТ2 получали 877 из "
          "13 026 пациентов (6,7%); в заранее заданной подгруппе результат согласуется с общим (слайд "
          "FIDELITY). Группа небольшая; комбинацию двух препаратов изучали отдельно в исследовании "
          "CONFIDENCE.")


# ----------------------------------------------------------------------------- slides 20 and 21 (shared)
COL_X, COL_W = 412, 512          # right column (after the chart is narrowed)
ND = "\u00a0"                      # no-break space for '95% ДИ 0,73–0,93'


def potassium_block(slide, y, facts, name="калий"):
    """Plate 'Калий' (#E97132) + grey card with three facts and the rule. Returns the bottom y."""
    paras = [[("Гиперкалиемия по сообщениям исследователей:", True), (" " + facts[0], False),
              (" Окончательная отмена из-за неё: " + facts[1], False),
              (" Смертельных исходов гиперкалиемии не было.", False)],
             [("Контроль:", True), (" до начала, через 4 недели после начала или повышения дозы, "
                                    "далее периодически.", False)]]
    plate(slide, COL_X, y, COL_W, 24, ORANGE, "Калий", f"Полоса_{name}")
    tb, h = flow(slide, COL_X + 10, y + 24 + 2 + 6, COL_W - 20, paras, f"Текст_{name}")
    grey_card(slide, COL_X, y + 26, COL_W, h + 12, f"Карточка_{name}")
    # card below the text in z-order is not needed: the text box is created first, so lift the card
    sp = slide.shapes._spTree
    card_el = one(slide, f"Карточка_{name}")._element
    sp.remove(card_el)
    sp.insert(sp.index(tb._element), card_el)
    return y + 26 + h + 12


def old_right_column(slide, names):
    for nm in names:
        delete_shape(one(slide, nm))


# ----------------------------------------------------------------------------- slide 20
def s20(prs):
    s = L.S(prs, 20)
    L.set_title(s, "FIDELIO-DKD: риск почечного исхода ниже на 18% по сравнению с плацебо")
    L.rep(s, "на фоне ингибитора АПФ или БРА", "на фоне ИАПФ или БРА")
    L.rep(s, "(5734 рандомизированы). Хроническая болезнь почек и сахарный диабет 2 типа; преимущественно "
             "3–4-я стадия болезни почек с выраженной альбуминурией.",
          "(5734 рандомизированы): ХБП и диабет 2 типа, преимущественно 3–4-я стадия с выраженной "
          "альбуминурией.")

    # chart: left edge on the text margin, plot narrowed to make room for the right column
    chart_transform(s, "", (138, 417), (68, 417), kx=318 / 408)
    style_curves(s)
    geom(one(s, "legend_finerenone"), x=82)
    geom(one(s, "legend_finerenone_label"), x=108)
    geom(one(s, "legend_placebo"), x=198)
    geom(one(s, "legend_placebo_label"), x=224)
    geom(one(s, "x_axis_title"), x=68, w=318)
    geom(one(s, "y_axis_title"), w=360)
    txt(s, 36, 466, 372, 26,
        [[("Кривая показывает накопленную частоту к каждому месяцу с учётом разного срока наблюдения, "
           "поэтому её конец выше доли за всё исследование.", 10, False, BLACK)]],
        "Сноска_график", spacing=1.0)

    old_right_column(s, ["event_fractions_label", "finerenone_event_rate", "placebo_event_rate",
                         "finerenone_group", "placebo_group", "hazard_ratio", "nnt", "hyperkalemia"])

    # block 'Результат'
    y = 168
    plate(s, COL_X, y, COL_W, 24, PURPLE, "Результат", "Полоса_результат")
    ix, iw = COL_X + 10, COL_W - 20
    yy = y + 26 + 6
    lab, hl = flow(s, ix, yy, iw, [[("Доля пациентов с событием за всё исследование", False)]], "Подпись_доли")
    yy += hl + 2
    numbers_row(s, ix, yy, iw, "17,8%", "21,1%", "Доля")
    yy += 36
    paras = [[("Риск ниже на 18% по сравнению с плацебо:", True),
              (f" отношение рисков 0,82 (95%{ND}ДИ{ND}0,73–0,93), p=0,001.", False)],
             [("ЧБНЛ:", True), (" чтобы предотвратить 1 событие за 3 года, лечат 29 пациентов.", False)],
             [("Сердечно-сосудистый исход (вторичный):", True),
              (f" риск ниже на 14% по сравнению с плацебо; 13,0% на финереноне и 14,8% на плацебо, "
               f"отношение рисков 0,86 (95%{ND}ДИ{ND}0,75–0,99).", False)]]
    tb, h = flow(s, ix, yy, iw, paras, "Текст_результат")
    bottom = yy + h + 6
    grey_card(s, COL_X, y + 26, COL_W, bottom - (y + 26), "Карточка_результат")
    sp = s.shapes._spTree
    c_el = one(s, "Карточка_результат")._element
    sp.remove(c_el)
    sp.insert(sp.index(lab._element), c_el)
    end = potassium_block(s, bottom + 8, ("18,3% на финереноне, 9,0% на плацебо.", "2,3% и 0,9%."))
    print("  s20 right column ends at", round(end))

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; ИАПФ: ингибитор "
            "ангиотензинпревращающего фермента; БРА: блокатор рецепторов ангиотензина II; ДИ: доверительный "
            "интервал.",
            ["ЧБНЛ: число пациентов, которых нужно лечить, чтобы предотвратить 1 событие. Гиперкалиемия: "
             "повышение калия сыворотки. Сердечно-сосудистый исход: сердечно-сосудистая смерть, нефатальные "
             "инфаркт миокарда и инсульт, госпитализация по поводу сердечной недостаточности.",
             "Bakris GL, et al. N Engl J Med. 2020;383:2219–2229."],
            abbr_lines=1, abbr_w=858)

    notes(s,
          "Сначала почки: исследование FIDELIO-DKD.",
          "В FIDELIO-DKD изучали взрослых с хронической болезнью почек и диабетом 2 типа. Финеренон "
          "сравнивали с плацебо на фоне ингибитора АПФ или блокатора рецепторов ангиотензина в максимально "
          "переносимой дозе. Первичный исход объединял почечную недостаточность, снижение рСКФ не менее чем "
          "на 40%, сохранявшееся не менее четырёх недель, и смерть от почечных причин.",
          "Группа плацебо показывает остаточный риск: на ИАПФ или БРА в максимально переносимой дозе "
          "почечный исход возник у 21,1% пациентов при медиане наблюдения 2,6 года. На финереноне у 17,8%.",
          "При медиане наблюдения 2,6 года отношение рисков составило 0,82, что соответствует относительному "
          "снижению на 18%. Доверительный интервал составил 0,73–0,93. За всё исследование исход "
          "зарегистрировали у 17,8% пациентов на финереноне и у 21,1% на плацебо. Кривая показывает "
          "накопленную частоту к каждому месяцу с учётом разного срока наблюдения, поэтому её конец не "
          "совпадает с долей за всё исследование.",
          "Исследователи сообщали о гиперкалиемии у 18,3% пациентов на финереноне и у 9,0% на плацебо. "
          "Окончательную отмену из-за неё зарегистрировали у 2,3% и 0,9% соответственно. Калий и рСКФ "
          "проверяют до начала лечения, через 4 недели после начала или повышения дозы и далее "
          "периодически; в исследовании гиперкалиемии со смертельным исходом не зарегистрировали.",
          "Почки показаны. Теперь сердце: FIGARO-DKD, следующий слайд.",
          "Фраза для визита: «В FIDELIO-DKD у пациентов с ХБП, диабетом 2 типа и альбуминурией на ИАПФ или "
          "БРА в максимально переносимой дозе финеренон снизил риск почечного исхода на 18% по сравнению с "
          "плацебо: исход возник у 17,8% против 21,1% при медиане наблюдения 2,6 года».",
          "Если врач спросит: «Есть ли польза для сердца?» Да, сердечно-сосудистый исход (вторичный): 13,0% "
          "против 14,8%, отношение рисков 0,86 (0,75–0,99). Подробно в FIGARO-DKD. «Сколько лечить?» ЧБНЛ 29 "
          "за 3 года. «Почему на графике 34% и 37%, а в тексте 17,8% и 21,1%?» Кривая показывает накопленную "
          "частоту к каждому месяцу, к концу графика наблюдаемых пациентов остаётся мало; 17,8% и 21,1% "
          "доля за всё исследование. «Почему в публикациях разные цифры по калию?» Исследователи сообщали о "
          "гиперкалиемии у 18,3% и 9,0%; по лабораторному критерию (калий выше 5,5 ммоль/л) в отдельном "
          "анализе было 21,4% и 9,2%, выше 6,0 ммоль/л 4,5% и 1,4% (Agarwal, JASN 2022). Определения "
          "разные, риск возрастал примерно вдвое, при регулярном контроле калия последствия сводились к "
          "минимуму.")


# ----------------------------------------------------------------------------- slide 21
def s21(prs):
    s = L.S(prs, 21)
    L.set_title(s, "FIGARO-DKD: риск сердечно-сосудистых событий ниже на 13% по сравнению с плацебо, основной "
                   "вклад дали госпитализации по поводу сердечной недостаточности (ниже на 29%)")
    L.rep(s, "на фоне ингибитора АПФ или БРА", "на фоне ИАПФ или БРА")
    # two-line title: the study text and the primary outcome go down by 12 and 6 pt
    shift([one(s, "study_population")], dy=12)
    shift([one(s, "primary_endpoint")], dy=6)

    chart_transform(s, "", (134, 417), (68, 417), kx=318 / 412)
    style_curves(s)
    geom(one(s, "legend_finerenone"), x=82)
    geom(one(s, "legend_finerenone_label"), x=108)
    geom(one(s, "legend_placebo"), x=198)
    geom(one(s, "legend_placebo_label"), x=224)
    geom(one(s, "x_axis_title"), x=68, w=318)
    geom(one(s, "y_axis_title"), w=360)

    old_right_column(s, ["event_fractions_label", "finerenone_event_rate", "placebo_event_rate",
                         "finerenone_group", "placebo_group", "hazard_ratio", "nnt", "heart_failure_result",
                         "hyperkalemia"])

    y = 168
    plate(s, COL_X, y, COL_W, 24, PURPLE, "Результат", "Полоса_результат")
    ix, iw = COL_X + 10, COL_W - 20
    yy = y + 26 + 6
    lab, hl = flow(s, ix, yy, iw, [[("Доля пациентов с событием за всё исследование", False)]], "Подпись_доли")
    yy += hl + 2
    numbers_row(s, ix, yy, iw, "12,4%", "14,2%", "Доля")
    yy += 36
    paras = [[("Риск ниже на 13% по сравнению с плацебо:", True),
              (f" отношение рисков 0,87 (95%{ND}ДИ{ND}0,76–0,98), p=0,03.", False)],
             [("ЧБНЛ:", True), (" чтобы предотвратить 1 событие за 3,5 года, лечат 47 пациентов.", False)]]
    tb, h = flow(s, ix, yy, iw, paras, "Текст_результат")
    yy += h + 6
    # component of the primary outcome: heart-failure hospitalisations
    big = txt(s, ix, yy, 78, 34, [[("−29%", 28, True, PURPLE)]], "Госпитализации_по_СН_число", spacing=1.0)
    hf = [[("Основной вклад дали госпитализации по поводу сердечной недостаточности (компонент первичного "
            f"исхода): риск ниже на 29%, отношение рисков 0,71 (95%{ND}ДИ{ND}0,56–0,90).", True)]]
    tb2, h2 = flow(s, ix + 80, yy, iw - 80, hf, "Госпитализации_по_СН_текст", size=16)
    bottom = yy + max(34, h2) + 6
    grey_card(s, COL_X, y + 26, COL_W, bottom - (y + 26), "Карточка_результат")
    sp = s.shapes._spTree
    c_el = one(s, "Карточка_результат")._element
    sp.remove(c_el)
    sp.insert(sp.index(lab._element), c_el)
    end = potassium_block(s, bottom + 8, ("10,8% на финереноне, 5,3% на плацебо.", "1,2% и 0,4%."))
    print("  s21 right column ends at", round(end))

    footers(s,
            "ХБП: хроническая болезнь почек; СН: сердечная недостаточность; ЧБНЛ: число пациентов, которых "
            "нужно лечить, чтобы предотвратить 1 событие; ИАПФ: ингибитор ангиотензинпревращающего фермента; "
            "БРА: блокатор рецепторов ангиотензина II; ДИ: доверительный интервал.",
            ["Pitt B, et al. N Engl J Med. 2021;385:2252–2263."])

    notes(s,
          "Почки мы видели в FIDELIO-DKD. Теперь сердце: исследование FIGARO-DKD, в нём группа пациентов "
          "шире, включая начальные стадии ХБП.",
          "В FIGARO-DKD участвовали взрослые с хронической болезнью почек и диабетом 2 типа. Финеренон "
          "сравнивали с плацебо на фоне ингибитора АПФ или блокатора рецепторов ангиотензина в максимально "
          "переносимой дозе. Первичный исход объединял сердечно-сосудистую смерть, нефатальные инфаркт и "
          "инсульт, а также госпитализацию по поводу сердечной недостаточности.",
          "При медиане наблюдения 3,4 года отношение рисков составило 0,87, что соответствует относительному "
          "снижению на 13%. Доверительный интервал составил 0,76–0,98. Основной вклад дали госпитализации по "
          "поводу сердечной недостаточности: риск ниже на 29% (отношение рисков 0,71; 95% ДИ 0,56–0,90). "
          "Снижение риска госпитализации по поводу сердечной недостаточности входит в показание у взрослых с "
          "ХБП при диабете 2 типа.",
          "Исследователи сообщали о гиперкалиемии у 10,8% пациентов на финереноне и у 5,3% на плацебо. "
          "Окончательную отмену из-за неё зарегистрировали у 1,2% и 0,4% соответственно. Калий и рСКФ "
          "проверяют до начала лечения, через 4 недели после начала или повышения дозы и далее периодически.",
          "FIDELIO-DKD и FIGARO-DKD вместе: объединённый анализ FIDELITY.",
          "Фраза для визита: «В FIGARO-DKD у пациентов с ХБП и диабетом 2 типа на ИАПФ или БРА финеренон "
          "снизил риск сердечно-сосудистых событий на 13% по сравнению с плацебо: 12,4% против 14,2% при "
          "медиане наблюдения 3,4 года. Основной вклад дали госпитализации по поводу сердечной "
          "недостаточности: риск ниже на 29%».",
          "Если врач спросит: «А остальные компоненты?» По отдельности различия по сердечно-сосудистой "
          "смерти, инфаркту и инсульту незначимы; основной вклад дали госпитализации по сердечной "
          "недостаточности. «А почки в FIGARO-DKD?» Почечный исход (вторичный) 9,5% против 10,8%, отношение "
          "рисков 0,87 (0,76–1,01), значимости не достиг. Почечную пользу показывают FIDELIO-DKD и FIDELITY. "
          "«А если у пациента есть сердечная недостаточность?» Пациентов с хронической симптомной сердечной "
          "недостаточностью со сниженной фракцией выброса в исследование не включали; сердечная "
          "недостаточность в анамнезе была у 7,7% участников объединённого анализа. Показание касается "
          "снижения риска госпитализации по сердечной недостаточности у взрослых с ХБП и диабетом 2 типа; "
          "лечение сердечной недостаточности как отдельного состояния в показание не входит, такие вопросы "
          "мы передаём в медицинский отдел.")


# ----------------------------------------------------------------------------- slide 22
def recolor_by_value(sh, mapping):
    """Recolour runs whose colour is a key of `mapping` (e.g. old blue/orange -> course palette)."""
    for r in sh._element.iter(qn("a:r")):
        c = r.find(qn("a:rPr") + "/" + qn("a:solidFill") + "/" + qn("a:srgbClr"))
        if c is not None and c.get("val") in mapping:
            c.set("val", mapping[c.get("val")])


def s22(prs):
    s = L.S(prs, 22)
    L.set_title(s, "FIDELITY: у 13 026 пациентов риск сердечно-сосудистых событий ниже на 14%, а почечных "
                   "исходов ниже на 23% по сравнению с плацебо")
    study = one(s, "study_population")
    rich(study, [[("13 026 пациентов с ХБП и диабетом 2 типа из FIDELIO-DKD и FIGARO-DKD, медиана "
                   "наблюдения 3,0 года. Заранее запланированный объединённый анализ данных каждого пациента. "
                   "Результат согласуется в подгруппах, в том числе у 877 пациентов (6,7%), получавших "
                   "иНГЛТ2.", False)]])
    geom(study, y=66, h=30)
    delete_shape(one(s, "column_separator"))
    delete_shape(one(s, "context_and_safety"))

    LIMIT = 490
    # potassium block (full width) is laid out first: it sets the bottom of the two panels
    k_paras = [[("Гиперкалиемия по сообщениям исследователей:", True),
                (" 14,0% на финереноне, 6,9% на плацебо. Окончательная отмена из-за неё: 1,7% и 0,6%. "
                 "Смертельных исходов гиперкалиемии не было. ", False),
                ("Контроль:", True),
                (" до начала, через 4 недели после начала или повышения дозы, далее периодически.", False)]]
    k_h = lines_of(k_paras, 868, 14) * 14 * 1.15 * 0.9
    k_top = LIMIT - (24 + 2 + k_h + 10)
    plate(s, 36, k_top, 888, 24, ORANGE, "Калий", "Полоса_калий")
    grey_card(s, 36, k_top + 26, 888, k_h + 10, "Карточка_калий")
    txt(s, 46, k_top + 26 + 5, 868, k_h + 2, [[(t, 14, b, BLACK) for t, b in k_paras[0]]], "Текст_калий")

    y_plate, h_plate = 99, 22
    card_top = y_plate + h_plate + 2
    card_bottom = k_top - 5
    IW = 412
    def_h, ax_h = 32, 15

    def results(pre):
        if pre == "cv_":
            ps = [[("Доля с событием за всё исследование: ", False),
                   ("финеренон 12,7%", False, PURPLE), (", ", False), ("плацебо 14,4%", False, GREY),
                   (".", False)],
                  [("Риск ниже на 14% по сравнению с плацебо:", True),
                   (f" отношение рисков 0,86 (95%{ND}ДИ{ND}0,78–0,95), p=0,0018. ЧБНЛ за 3 года: 46 "
                    f"(95%{ND}ДИ{ND}29–109).", False)],
                  [("Из компонентов:", True),
                   (f" госпитализации по поводу сердечной недостаточности, риск ниже на 22% (отношение "
                    f"рисков 0,78; 95%{ND}ДИ{ND}0,66–0,92).", False)]]
        else:
            ps = [[("Доля с событием за всё исследование: ", False),
                   ("финеренон 5,5%", False, PURPLE), (", ", False), ("плацебо 7,1%", False, GREY),
                   (".", False)],
                  [("Риск ниже на 23% по сравнению с плацебо:", True),
                   (f" отношение рисков 0,77 (95%{ND}ДИ{ND}0,67–0,88), p=0,0002. ЧБНЛ за 3 года: 60 "
                    f"(95%{ND}ДИ{ND}38–142).", False)]]
        return ps

    cv_res = results("cv_")
    res_h = lines_of([[(r[0], r[1]) for r in p] for p in cv_res], IW, 14) * 14 * 1.15 * 0.9 + 2 * (len(cv_res) - 1)
    plot_h = card_bottom - 4 - res_h - (card_top + 5 + def_h + 2) - ax_h - 9.5 - 47
    print("  s22 plot height", round(plot_h, 1), "results", round(res_h), "potassium block from", round(k_top))
    ky = plot_h / 129

    panels = [
        ("cv_", 36, 140, "Комбинированный сердечно-сосудистый исход", "cv_heading", "cv_definition",
         [[("Сердечно-сосудистая смерть, нефатальный инфаркт, нефатальный инсульт или госпитализация по "
            "поводу сердечной недостаточности.", False)]]),
        ("renal_", 492, 602, "Комбинированный почечный исход", "renal_heading", "renal_definition",
         [[("Почечная недостаточность, снижение рСКФ на ", False), ("≥57%", True),
           (", сохраняющееся не менее 4 недель, или смерть от почечных причин.", False)]]),
    ]
    for pre, PL, ox, head, head_nm, def_nm, def_paras in panels:
        IL = PL + 10
        delete_shape(one(s, head_nm))
        plate(s, PL, y_plate, 432, h_plate, PURPLE, head, f"Полоса_{pre}результат")
        grey_card(s, PL, card_top, 432, card_bottom - card_top, f"Карточка_{pre}результат")
        sp = s.shapes._spTree
        el = one(s, f"Карточка_{pre}результат")._element
        sp.remove(el)
        sp.insert(2, el)               # the card goes behind the chart and the texts
        d = one(s, def_nm)
        rich(d, def_paras)
        geom(d, x=IL, y=card_top + 5, w=IW, h=def_h)
        y_ax = card_top + 5 + def_h + 2
        geom(one(s, pre + "y_axis_title"), x=IL, y=y_ax, w=300, h=ax_h)
        ny = y_ax + ax_h + 9.5 + plot_h
        chart_transform(s, pre, (ox, 332), (PL + 42, ny), kx=360 / 312, ky=ky)
        style_curves(s, pre)
        geom(one(s, pre + "legend_finerenone"), x=PL + 50, y=y_ax + ax_h + 21)
        geom(one(s, pre + "legend_finerenone_label"), x=PL + 72, y=y_ax + ax_h + 11)
        geom(one(s, pre + "legend_placebo"), x=PL + 180, y=y_ax + ax_h + 21)
        geom(one(s, pre + "legend_placebo_label"), x=PL + 204, y=y_ax + ax_h + 11)
        geom(one(s, pre + "x_axis_title"), x=PL + 42, y=ny + 31, w=360, h=16)
        for nm in ("event_fractions", "hazard_ratio", "nnt"):
            delete_shape(one(s, pre + nm))
        flow(s, IL, ny + 47, IW, results(pre), f"{pre}результаты", spacing=0.9, after=2)

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; ДИ: "
            "доверительный интервал; ЧБНЛ: число пациентов, которых нужно лечить в течение 3 лет, чтобы "
            "предотвратить 1 событие; иНГЛТ2: ингибитор натрий-глюкозного котранспортёра 2 типа; СН: "
            "сердечная недостаточность.",
            ["Agarwal R, et al. Eur Heart J. 2022;43:474–484. p-значения поискового анализа без поправки на "
             "множественные сравнения."])

    notes(s,
          "В каждом исследовании мы видели свою сторону: почки в FIDELIO-DKD и сердце в FIGARO-DKD. FIDELITY "
          "объединяет их данные.",
          "FIDELITY объединил данные FIDELIO-DKD и FIGARO-DKD: 13 026 пациентов с хронической болезнью почек "
          "и диабетом 2 типа, медиана наблюдения три года. При сравнении финеренона с плацебо относительный "
          "риск комбинированного сердечно-сосудистого исхода был ниже на 14%. Для комбинированного "
          "почечного исхода относительное снижение риска составило 23%. Почечный исход включал почечную "
          "недостаточность, снижение рСКФ не менее чем на 57%, сохранявшееся не менее четырёх недель, и "
          "смерть от почечных причин.",
          "Исследователи сообщали о гиперкалиемии у 14,0% пациентов на финереноне и у 6,9% на плацебо. "
          "Окончательную отмену из-за неё зарегистрировали у 1,7% и 0,6% соответственно. Калий и рСКФ "
          "проверяют до начала лечения, через 4 недели после начала или повышения дозы и далее периодически.",
          "Одним из первых признаков ответа служит альбуминурия. Что с ней происходит в первые месяцы, на "
          "следующем слайде.",
          "Фраза для визита: «В объединённом анализе FIDELITY у 13 026 пациентов с ХБП и диабетом 2 типа при "
          "медиане наблюдения 3 года финеренон по сравнению с плацебо снизил риск сердечно-сосудистых событий "
          "на 14%, риск почечных исходов на 23% и риск госпитализаций по поводу сердечной недостаточности на "
          "22%. Результат согласуется и у пациентов, получавших иНГЛТ2».",
          "Если врач спросит: «А пациенты на иНГЛТ2?» На старте иНГЛТ2 получали 877 пациентов (6,7%), "
          "арГПП-1 944 (7,2%); в заранее заданных подгруппах результат согласуется с общим. "
          "Рандомизированного сравнения комбинации в FIDELITY не было; об отдельном исследовании комбинации "
          "см. CONFIDENCE. «Что с калием?» Госпитализация из-за тяжёлой гиперкалиемии 0,9% против 0,2%, "
          "гипокалиемия реже: 1,1% против 2,3%. «Насколько надёжен объединённый анализ?» Анализ спланирован "
          "заранее, до объединения данных; его тесты поисковые, p без поправки на множественные сравнения, "
          "поэтому опираемся на отношения рисков и доверительные интервалы.")


# ----------------------------------------------------------------------------- slide 23
def s23(prs):
    s = L.S(prs, 23)
    L.set_title(s, "К 4-му месяцу АКО на финереноне на 32% ниже, чем на плацебо")

    # top block: one line
    top = one(s, "population")
    rich(top, [[("Кто участвовал: ", True), ("13 026 пациентов с ХБП и диабетом 2 типа. ", False),
                ("Что сравнивали: ", True),
                ("финеренон или плацебо на фоне ИАПФ или БРА в максимально переносимой дозе.", False)]])
    geom(top, x=36, y=62, w=888, h=20)
    delete_shape(one(s, "background_therapy"))
    delete_shape(one(s, "marker_definition"))
    delete_shape(one(s, "monitoring"))

    # left column: plate + grey card
    delete_shape(one(s, "left_heading"))
    plate(s, 36, 102, 345, 48, TEAL, "Относительная разница с плацебо\nпо АКО к 4-му месяцу",
          "Полоса_разница")
    grey_card(s, 36, 154, 345, 150, "Карточка_разница")
    sp = s.shapes._spTree
    el = one(s, "Карточка_разница")._element
    sp.remove(el)
    sp.insert(sp.index(one(s, "placebo_adjusted_uacr_reduction")._element), el)
    geom(one(s, "placebo_adjusted_uacr_reduction"), x=46, y=160, w=325, h=34)
    geom(one(s, "mean_change_explanation"), x=46, y=196, w=325, h=20)
    ratio = one(s, "ratio_detail")
    geom(ratio, x=46, y=218, w=325, h=30)
    extra, h_extra = flow(s, 46, 258, 325,
                          [[("В отдельных исследованиях АКО было ниже, чем на плацебо, на 31% (FIDELIO-DKD) и "
                             "на 32% (FIGARO-DKD).", False)]], "Отдельные_исследования")
    print("  s23 card text ends at", round(258 + h_extra), "(card bottom 304)")

    # chart and the line under it
    ch = one(s, "Chart")
    geom(ch, y=100, h=232)
    # chart title: the line break sat inside one run; make it two paragraphs (renders the same everywhere)
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    for r in ch.chart._chartSpace.iter("{%s}r" % A):
        t = r.find("{%s}t" % A)
        if t is not None and "\n" in (t.text or "") and r.getparent().getparent().getparent().tag == "{%s}rich" % C:
            p0 = r.getparent()
            first, second = t.text.split("\n", 1)
            t.text = first
            p1 = copy.deepcopy(p0)
            p1.find("{%s}r" % A).find("{%s}t" % A).text = second
            p0.addnext(p1)
            break
    sec = one(s, "secondary_analysis_label")
    rich(sec, [[("Доля пациентов, у которых АКО снизилось на ≥30% от исходного: ", True),
                ("финеренон 53,2%, плацебо 27,0% (дополнительный анализ, не запланированный заранее; "
                 "12 512 пациентов).", False)]])
    geom(sec, x=418, y=336, w=506, h=44)

    # conclusion (16 pt): main phrase bold purple, second sentence plain
    concl = one(s, "interpretation")
    rich(concl, [[("В дополнительном анализе снижение АКО к 4-му месяцу объясняло большую часть эффекта "
                   "финеренона на почки: от 64% до 84% в зависимости от метода анализа.", True)]], color=PURPLE)
    geom(concl, x=36, y=394, w=888, h=36)
    lim = one(s, "individual_limit")
    rich(lim, [[("АКО помогает оценивать ответ на терапию.", False)]])
    geom(lim, x=36, y=436, w=888, h=22)

    footers(s,
            "АКО: отношение альбумина к креатинину мочи; ХБП: хроническая болезнь почек; ИАПФ: ингибитор "
            "ангиотензинпревращающего фермента; БРА: блокатор рецепторов ангиотензина II; ДИ: доверительный "
            "интервал.",
            ["Финеренон-Орвилле. Инструкция по медицинскому применению. Agarwal R, et al. Eur Heart J. "
             "2022;43:474–484. Agarwal R, et al. Ann Intern Med. 2023;176:1606–1616. KDIGO 2024. Kidney Int. "
             "105(4S):S117–S314."])

    notes(s,
          "Мы видели исходы. Теперь вернёмся к первому признаку ответа: альбуминурии в первые месяцы "
          "лечения.",
          "В FIDELITY к четвёртому месяцу отношение альбумина к креатинину мочи при финереноне было на 32% "
          "ниже, чем при плацебо, с поправкой на исходный уровень. Разница между группами относительная. На "
          "слайде отдельно показана доля пациентов, у которых АКО снизилось не менее чем на 30% от "
          "собственного исходного уровня: 53,2% при финереноне и 27,0% при плацебо. Этот показатель "
          "рассчитали в дополнительном анализе после получения данных исследования. По инструкции в "
          "FIDELIO-DKD АКО к 4-му месяцу было ниже на 31%, в FIGARO-DKD на 32% по сравнению с плацебо.",
          "Раннее изменение альбуминурии помогает оценивать ответ на лечение. В дополнительном анализе "
          "снижение АКО объясняло от 64% до 84% почечного эффекта и от 26% до 37% сердечно-сосудистого.",
          "Альбуминурию снижает и ингибитор НГЛТ2. Часто возникает вопрос, что будет, если начать оба "
          "препарата. Об этом исследование CONFIDENCE.",
          "Фраза для визита: «В FIDELITY к 4-му месяцу АКО на финереноне было на 32% ниже, чем на плацебо. В "
          "дополнительном анализе это раннее снижение объясняло большую часть почечного эффекта "
          "финеренона».",
          "Если врач спросит: «Можно ли по АКО предсказать исход у конкретного пациента?» АКО показывает "
          "ответ на уровне групп; индивидуальный прогноз врач оценивает отдельно, с учётом калия и функции "
          "почек, которые продолжают контролировать.")


# ----------------------------------------------------------------------------- slide 24
def number_format(chart, code):
    """Series data labels: number format with the Russian locale tag (decimal comma)."""
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    for dl in chart._chartSpace.iter("{%s}dLbls" % C):
        if dl.getparent().tag != "{%s}ser" % C:
            continue
        for k in dl.findall("{%s}numFmt" % C):
            dl.remove(k)
        nf = etree.Element("{%s}numFmt" % C)
        nf.set("formatCode", code)
        nf.set("sourceLinked", "0")
        dl.insert(0, nf)


def s24(prs):
    s = L.S(prs, 24)
    L.set_title(s, "CONFIDENCE: при одновременном старте финеренона и эмпаглифлозина АКО снизилось сильнее, "
                   "чем на каждом препарате отдельно")

    # study description: plate + grey card with three lines
    pat = one(s, "Пациенты")
    plate(s, 36, 70, 888, 22, TEAL, "Кто участвовал и что сравнивали", "Полоса_исследование")
    grey_card(s, 36, 94, 888, 58, "Карточка_исследование")
    sp = s.shapes._spTree
    el = one(s, "Карточка_исследование")._element
    sp.remove(el)
    sp.insert(sp.index(pat._element), el)
    rich(pat, [[("800 взрослых с ХБП и диабетом 2 типа. Двойное слепое исследование, 180 дней.", False)],
               [("рСКФ 30–90 мл/мин/1,73 м²; АКО 100–5000 мг/г; калий ≤4,8 ммоль/л; ИАПФ или БРА.", False)],
               [("Что сравнивали: ", True), ("финеренон и эмпаглифлозин одновременно, финеренон отдельно, "
                                             "эмпаглифлозин отдельно.", False)]], size=15)
    geom(pat, x=46, y=99, w=868, h=50)
    set_line_spacing(pat, 90)

    # chart headings and units
    for nm, y in (("АКО заголовок", 158), ("Калий заголовок", 158)):
        sh = one(s, nm)
        recolor_text(sh, BLACK, bold=True, size=18)
        geom(sh, y=y, h=24)
    L.rep(s, "Отношение к исходному уровню, оценка по модели",
          "Отношение к исходному уровню (1,0 = исходный); в скобках 95% ДИ")
    geom(one(s, "АКО единицы"), y=184, h=19)
    geom(one(s, "Калий единицы"), y=184, h=19)
    c1, c2 = one(s, "АКО на 180 день"), one(s, "Калий выше 5,5")
    geom(c1, y=204, h=166)
    geom(c2, y=204, h=166)
    number_format(c1.chart, "[$-419]0.00")
    number_format(c2.chart, '[$-419]0.0"%"')

    # two lines under each column of the left chart: ratio with CI and the change from baseline
    delete_shape(one(s, "АКО ДИ"))
    for cx, ratio, change, nm in ((112, "0,48 (0,44–0,54)", "−52% от исходного", "комбинация"),
                                  (248, "0,68 (0,61–0,76)", "−32% от исходного", "финеренон"),
                                  (384, "0,71 (0,64–0,79)", "−29% от исходного", "эмпаглифлозин")):
        t = txt(s, cx - 66, 372, 132, 30, [[(ratio, 14, False, BLACK)], [(change, 14, True, BLACK)]],
                f"Подпись_столбца_{nm}", spacing=0.9)
        for p in t.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
    k = one(s, "Калий сравнение")
    rich(k, [[("Гиперкалиемия (калий выше 5,5 ммоль/л хотя бы раз): ", True),
              ("15,1% на комбинации, 18,8% на финереноне; различие незначимо (p=0,85).", False)]])
    geom(k, x=500, y=372, w=424, h=30)

    # conclusion and the card with the primary endpoint and the potassium rule
    res = one(s, "Результат")
    rich(res, [[("Комбинация снизила АКО сильнее каждого препарата отдельно: ниже, чем на финереноне, на 29%, "
                 "и ниже, чем на эмпаглифлозине, на 32% (p<0,001 для обоих сравнений).", True)]], color=PURPLE)
    geom(res, x=36, y=410, w=888, h=36)
    lim = one(s, "Граница результата")
    grey_card(s, 36, 448, 888, 40, "Карточка_первичная_точка")
    sp = s.shapes._spTree
    el = one(s, "Карточка_первичная_точка")._element
    sp.remove(el)
    sp.insert(sp.index(lim._element), el)
    rich(lim, [[("Первичная точка:", True), (" изменение АКО от исходного до 180-го дня. ", False),
                ("Калий и рСКФ проверяют до начала лечения, через 4 недели после начала или повышения дозы и "
                 "далее периодически.", False)]])
    geom(lim, x=46, y=453, w=868, h=30)

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; АКО: отношение "
            "альбумина к креатинину мочи; ИАПФ: ингибитор ангиотензинпревращающего фермента; БРА: блокатор "
            "рецепторов ангиотензина II; ДИ: доверительный интервал.",
            [L.shape_named(s, "Источники").text_frame.text])

    notes(s,
          "Альбуминурию снижает и ингибитор НГЛТ2. Часто возникает вопрос, как финеренон сочетается с ним. "
          "На эту тему есть исследование CONFIDENCE.",
          "В CONFIDENCE сравнивали одновременное начало приёма финеренона и эмпаглифлозина с началом приёма "
          "каждого препарата отдельно. В основной анализ вошли 800 взрослых с хронической болезнью почек и "
          "диабетом 2 типа. Все получали ингибитор АПФ или блокатор рецепторов ангиотензина в максимально "
          "переносимой дозе.",
          "Слева показано отношение АКО к исходному уровню на 180-й день: 0,48 при комбинации, 0,68 при "
          "финереноне и 0,71 при эмпаглифлозине. Такие отношения соответствуют снижению внутри групп на 52%, "
          "32% и 29%. В подписях приведены 95% доверительные интервалы. При сравнении групп АКО на "
          "комбинации было на 29% ниже, чем на финереноне, и на 32% ниже, чем на эмпаглифлозине. В обоих "
          "сравнениях значение p было меньше 0,001. Результаты относятся к альбуминурии. Исходы в "
          "исследовании не оценивали; пользу по исходам показывают FIDELIO-DKD, FIGARO-DKD и FIDELITY.",
          "Справа показана доля пациентов, у которых калий хотя бы раз превысил 5,5 ммоль/л: 15,1% при "
          "комбинации, 18,8% при финереноне и 9,7% при эмпаглифлозине. Различие между комбинацией и "
          "финереноном по шансам гиперкалиемии незначимо (p=0,85). Калий при сочетании контролируют по "
          "тому же правилу.",
          "Что это значит для выбора препарата: финеренон и иНГЛТ2 дополняют друг друга, врачу не нужно "
          "выбирать между ними. Об этом следующий слайд.",
          "Фраза для визита: «В исследовании CONFIDENCE у 800 пациентов с ХБП и диабетом 2 типа "
          "одновременный старт финеренона и эмпаглифлозина снизил АКО к 180-му дню на 52% от исходного. На "
          "одном финереноне снижение составило 32%, на одном эмпаглифлозине 29%; различие комбинации с "
          "каждым из препаратов значимо (p<0,001). Гиперкалиемия при комбинации возникла у 15,1% пациентов, "
          "на одном финереноне у 18,8%».",
          "Если врач спросит: «Есть ли данные по исходам у комбинации?» В CONFIDENCE оценивали АКО за 180 "
          "дней, исходы не оценивали. Польза финеренона по исходам показана в FIDELIO-DKD, FIGARO-DKD и "
          "FIDELITY, в том числе у пациентов на иНГЛТ2. «Эмпаглифлозин снижает риск гиперкалиемии?» В "
          "CONFIDENCE при добавлении эмпаглифлозина значимого снижения гиперкалиемии не было: 15,1% на "
          "комбинации и 18,8% на финереноне (p=0,85); на одном эмпаглифлозине 9,7%. «Насколько комбинация "
          "лучше?» АКО на комбинации ниже, чем на финереноне, на 29% (отношение 0,71; 95% ДИ 0,61–0,82), и "
          "ниже, чем на эмпаглифлозине, на 32% (0,68; 0,59–0,79); это сравнение между группами, не путать со "
          "снижением от исходного (52%, 32%, 29%). «Почему нет группы плацебо?» В каждой группе был "
          "активный препарат и плацебо-пара другого препарата; чистой группы плацебо в исследовании не "
          "было. «Можно ли начинать финеренон и эмпаглифлозин одновременно?» В CONFIDENCE изучали "
          "одновременный старт. Схема одновременного старта двух препаратов в инструкции не описана: "
          "медицинский представитель пересказывает результаты исследования и не даёт рекомендаций по схеме "
          "назначения; порядок назначения определяет врач.")


# ----------------------------------------------------------------------------- run
def run(prs):
    s17(prs)
    s18(prs)
    s19(prs)
    s20(prs)
    s21(prs)
    s22(prs)
    s23(prs)
    s24(prs)
