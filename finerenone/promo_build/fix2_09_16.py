"""Fix2 for original slides 9-16: corrections after checking v1 (applied after impl_*, before impl_final).
Addresses original slide numbers (base47): 13 -> v1 slide 14, 14 -> v1 slide 20, 15 -> 15, 16 -> 16.

0. s10: Alfego 2021 (customer decision: only in the notes) added as an answer block.
1. s13: notes start repeated the end of H2 ("Portrait") -> trimmed; line break before "из-за" (no "из-/за");
   footnote 1 states that the data are for semaglutide; bottom line 2 pt lower (clear of the membrane picture).
2. s14: notes start repeated the end of H3 -> reworded.
3. s15: category labels of the chart on one line (two-line labels were tight, overlap.py flagged them);
   right card aligned to the right edge 924 of the grey card and the plate.
4. s16: aldosterone markers in the green of the aldosterone label/arrow; pointer line black.
"""
import copy

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt

import r14lib as L
from lib import iter_shapes

GREEN = "196B24"


def _geom(sh, x=None, y=None, w=None, h=None):
    if x is not None:
        sh.left = Pt(x)
    if y is not None:
        sh.top = Pt(y)
    if w is not None:
        sh.width = Pt(w)
    if h is not None:
        sh.height = Pt(h)


def s10(prs):
    """Decision of the customer (analysis_promo_47, item 12): Alfego 2021 only in the notes of this slide.
    impl_09_16 did not add it."""
    s = L.S(prs, 10)
    blocks = [p.text for p in s.notes_slide.notes_text_frame.paragraphs if p.text.strip()]
    assert not any("Alfego" in b for b in blocks)
    blocks.append(
        "Если врач спросит: «Как часто АКО определяют на практике?» Ответ: «Это американские данные: среди "
        "28 295 982 пациентов с диабетом и/или гипертензией АКО определяли у 21,0%, а рСКФ у 89,6% "
        "(Alfego D, et al. Diabetes Care. 2021;44:2025–2032; лаборатория Labcorp, 2013–2019 годы). Они "
        "показывают разрыв между двумя анализами, но не описывают практику у нас.»")
    L.set_notes(s, "\n\n".join(blocks))


def s13(prs):
    s = L.S(prs, 13)
    # notes: H2 (the previous slide) already ends with "Следующий слайд показывает четыре группы препаратов"
    L.rep(s, "Пациента с ХБП и диабетом мы нашли. Чем его лечат сегодня? Для защиты почек",
          "Для защиты почек", count=1, notes=True)

    # "из-за" must not be torn apart: break the line before it
    sh = L.shape_with(s, "госпитализации из-за сердечной недостаточности")
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            t = r.text
            key = "госпитализации из-за сердечной недостаточности"
            if key in t:
                i = t.index(key) + len("госпитализации")
                head, tail = t[:i], t[i + 1:]          # drop the space before "из-за"
                r.text = head
                br = etree.SubElement(p._p, qn("a:br"))
                rpr = r._r.find(qn("a:rPr"))
                if rpr is not None:
                    br.append(copy.deepcopy(rpr))
                new_r = copy.deepcopy(r._r)
                new_r.find(qn("a:t")).text = tail
                r._r.addnext(br)
                br.addnext(new_r)
    # footnote 1: the data belong to semaglutide only
    L.rep(s, "¹ Семаглутид, исследование FLOW:", "¹ Данные семаглутида, исследование FLOW:", count=1)
    # bottom line: 2 pt lower, away from the tail of the membrane receptor picture
    b = L.shape_with(s, "Финеренон добавляют к ингибитору АПФ или БРА. Набор препаратов")
    b.top = Pt(464)


def s14(prs):
    s = L.S(prs, 14)
    # H3 (the previous slide) ends with "Сначала договоримся, как читать результат: ..."
    L.rep(s, "Мы показали, откуда берётся показание. Теперь доказательства, и сначала короткая памятка: как читать "
             "числа исследования и как пересказать их врачу. Чтобы понять",
          "Это короткая памятка: как читать числа исследования и как пересказать их врачу. Чтобы понять",
          count=1, notes=True)


def s15(prs):
    s = L.S(prs, 15)
    # chart: one-line category labels instead of two lines with tight spacing
    ch = [sh for sh in s.shapes if sh.has_chart][0].chart
    cd = CategoryChartData(number_format="0.0%")
    cd.categories = ["Плацебо (340 из 2199)", "Канаглифлозин (245 из 2202)"]
    cd.add_series("Первичный исход", (340 / 2199, 245 / 2202))
    ch.replace_data(cd)
    # right card up to the right edge 924 (as the grey card above and the plate below)
    card = L.shape_named(s, "Карточка_результат")
    card.width = Pt(924 - 603)
    for name in ("Отношение рисков", "Доверительный интервал", "Наблюдение", "Граница доказательств"):
        L.shape_named(s, name).width = Pt(291)


def s16(prs):
    s = L.S(prs, 16)
    for sh in iter_shapes(s.shapes):
        if sh.name in ("Условное обозначение альдостерона", "Альдостерон связан с рецептором"):
            for c in sh._element.spPr.iter(qn("a:srgbClr")):
                c.set("val", GREEN)
        if sh.name == "Указатель_насос":
            sh.line.color.rgb = RGBColor.from_string("000000")


def run(prs):
    s10(prs)
    s13(prs)
    s14(prs)
    s15(prs)
    s16(prs)
