"""53 -> 49 slides: remove repeats (final numbers 42, 43, 44, 50), fix the notes and the one slide
text that pointed to them. Pictures and visual slides are not touched."""
import copy
import os
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import r14lib as L

REMOVE = [42, 43, 44, 50]


def add_para_after(slide, needle, text):
    """Blank line and a new notes paragraph after the one that contains `needle`, same format."""
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn("a:p")):
        if needle in "".join(t.text or "" for t in p.iter(qn("a:t"))):
            new = copy.deepcopy(p)
            runs = new.findall(qn("a:r"))
            runs[0].find(qn("a:t")).text = text
            for r in runs[1:]:
                new.remove(r)
            for br in new.findall(qn("a:br")):
                new.remove(br)
            blank = copy.deepcopy(new)
            for r in blank.findall(qn("a:r")):
                blank.remove(r)
            p.addnext(new)
            p.addnext(blank)
            return
    raise KeyError(needle)


def edit(prs):
    S = lambda n: prs.slides[n - 1]
    # 41 (visit): the two opening questions are now answered in the self-check
    L.rep(S(41), "Вернёмся к двум вопросам из начала курса и ответим на них.", "Теперь проверим себя.", notes=True)
    # 45 (self-check): close the loop with the two situations from slide 3
    L.rep(S(45), "К ответам вернёмся после обсуждения каждого вопроса.",
          "К ответам вернёмся после обсуждения каждого вопроса. Вопросы 3 и 4 отвечают и на две ситуации "
          "из начала курса.", notes=True)
    # 49 (FINE-ONE): the FDA decision slide goes, its facts move here
    L.rep(S(49), "16.09.2026, подробности на следующем слайде.",
          "одобрено 16.09.2026 по снижению АКО и данным при диабете 2 типа.")
    L.rep(S(49), "БРА: блокатор рецепторов ангиотензина II; ДИ: доверительный интервал.",
          "БРА: блокатор рецепторов ангиотензина II; ДИ: доверительный интервал; FDA: Управление по контролю "
          "качества пищевых продуктов и лекарств США.")
    L.rep(S(49), "поэтому показание FDA сформулировано через ожидаемый эффект (следующий слайд).",
          "поэтому показание FDA сформулировано через ожидаемый эффект.", notes=True)
    L.rep(S(49), " Решение FDA рассматриваем на следующем слайде.", "", notes=True)
    add_para_after(S(49), "передавайте в медицинский отдел.",
                   "Решение FDA. 16 сентября 2026 года (объявлено 17 сентября) FDA одобрило дополнение к "
                   "американской инструкции Керендии для взрослых с ХБП, связанной с диабетом 1 типа. Показание "
                   "в США: снижение отношения альбумина к креатинину мочи, которое, как ожидается, уменьшит риск "
                   "устойчивого снижения рСКФ и терминальной стадии болезни почек. Основанием служили снижение "
                   "АКО в FINE-ONE (суррогатная конечная точка) и данные исследований при диабете 2 типа. "
                   "Ожидаемую пользу нельзя считать прямым доказательством снижения потребности в диализе: такие "
                   "исходы в FINE-ONE не оценивали. Решение действует в США и относится к препарату Керендия; в "
                   "инструкции Финеренон-Орвилле такого показания нет. Решений других регуляторов по сообщениям нет.")
    add_para_after(S(49), "Решение FDA. 16 сентября", "Дальше ХБП без диабета, исследование FIND-CKD.")


def remove(prs):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for n in sorted(REMOVE, reverse=True):
        el = ids[n - 1]
        prs.part.drop_rel(el.rId)
        lst.remove(el)


def main(src, dst):
    prs = Presentation(src)
    assert len(prs.slides) == 53
    edit(prs)
    remove(prs)
    assert len(prs.slides) == 49
    prs.save(dst)


if __name__ == "__main__":
    main(*sys.argv[1:])
