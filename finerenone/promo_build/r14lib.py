"""Helpers for round 14 (fixes from review_48_slides.md). Matching ignores the difference between
normal and non-breaking spaces; every edit raises if its target text is not found."""
import copy
import re

from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from lib import iter_shapes, set_paragraph_texts, shape_lines, set_rich_lines, delete_shape

NB = " "


def S(prs, n):
    """Slide by 1-based number."""
    return prs.slides[n - 1]


def _norm(s):
    return s.replace(NB, " ")


def rep(slide, old, new, count=None, notes=False):
    """Replace `old` by `new` inside paragraphs of a slide (or its notes), across runs, keeping the
    format of the run where the match starts. NBSP and space are treated as equal in matching.
    Returns the number of replacements; raises KeyError if none (or if count is given and differs)."""
    root = slide.notes_slide.notes_text_frame._txBody if notes else slide.shapes._spTree
    n_done = 0
    o = _norm(old)
    for p_el in root.iter(qn("a:p")):
        start = 0
        while True:
            runs = p_el.findall(qn("a:r"))
            texts = [(r.find(qn("a:t")).text or "") for r in runs]
            full = _norm("".join(texts))
            idx = full.find(o, start)
            if idx < 0:
                break
            end = idx + len(o)
            pos = 0
            s_i = e_i = None
            for i, t in enumerate(texts):
                if s_i is None and idx < pos + len(t):
                    s_i, s_off = i, idx - pos
                if end <= pos + len(t):
                    e_i, e_off = i, end - pos
                    break
                pos += len(t)
            ts = runs[s_i].find(qn("a:t"))
            if s_i == e_i:
                ts.text = texts[s_i][:s_off] + new + texts[s_i][e_off:]
            else:
                ts.text = texts[s_i][:s_off] + new
                runs[e_i].find(qn("a:t")).text = texts[e_i][e_off:]
                for r in runs[s_i + 1:e_i]:
                    p_el.remove(r)
            for r in p_el.findall(qn("a:r")):
                t = r.find(qn("a:t"))
                if t.text and t.text != t.text.strip():
                    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            start = idx + len(new)
            n_done += 1
    if n_done == 0 or (count is not None and n_done != count):
        raise KeyError(f"rep: found {n_done} of {old!r}")
    return n_done


def shape_with(slide, needle):
    """First shape (inside groups too) whose text contains `needle` (NBSP-insensitive)."""
    nd = _norm(needle)
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame and nd in _norm(sh.text_frame.text):
            return sh
    raise KeyError(needle)


def shape_named(slide, name):
    for sh in iter_shapes(slide.shapes):
        if sh.name == name:
            return sh
    raise KeyError(name)


def title_shape(slide):
    """Topmost text shape with a 20 pt run above y=90 pt."""
    best = None
    for sh in iter_shapes(slide.shapes):
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()]
        if sizes and max(sizes) >= 20 and Emu(sh.top).pt < 90:
            if best is None or sh.top < best.top:
                best = sh
    if best is None:
        raise KeyError("title")
    return best


def set_title(slide, text):
    """New title; ' / ' starts a new paragraph (only if the old title already had two paragraphs,
    otherwise pass a single line). Keeps the first run format of each paragraph."""
    sh = title_shape(slide)
    lines = text.split(" / ")
    set_paragraph_texts(sh, lines)
    return sh


def set_lines(shape, lines):
    """Rewrite paragraphs of a shape, keeping first-run format per paragraph."""
    set_paragraph_texts(shape, lines)


def set_notes(slide, text):
    """Replace speaker notes. Paragraphs separated by '\n'."""
    tf = slide.notes_slide.notes_text_frame
    tf.text = text


def notes_text(slide):
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def delete_paragraph_with(slide, needle):
    """Delete the paragraph that contains `needle` (NBSP-insensitive)."""
    nd = _norm(needle)
    for sh in iter_shapes(slide.shapes):
        if not sh.has_text_frame:
            continue
        for p in list(sh.text_frame.paragraphs):
            if nd in _norm("".join(r.text for r in p.runs)):
                p._p.getparent().remove(p._p)
                return sh
    raise KeyError(needle)
