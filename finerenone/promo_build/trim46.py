"""49 -> 46 slides: remove the CYP3A4 slide, the interactions slide and the summary (numbers 34, 35, 43
in the 49-slide deck). The interaction rules move into the notes of the potassium slide (33), the
notes junctions are fixed. Pictures on the remaining slides are not touched."""
import os
import sys

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import r14lib as L
from trim49 import add_para_after

REMOVE = [34, 35, 43]


def edit(prs):
    S = lambda n: prs.slides[n - 1]
    # 33 (potassium): interactions from the removed slides 34 and 35, then the way to the safety slide
    L.rep(S(33), "Её меняют лекарства, которые влияют на фермент CYP3A4. Об этом следующий слайд.",
          "Её меняют лекарства, которые влияют на фермент CYP3A4: около 90% финеренона превращается в "
          "неактивные метаболиты с его участием.", notes=True)
    add_para_after(S(33), "около 90% финеренона превращается",
                   "Взаимодействия. Сильные ингибиторы CYP3A4 противопоказаны: кларитромицин, кетоконазол, "
                   "итраконазол, ритонавир, нелфинавир. Концентрация финеренона при них повышается. Финеренон не "
                   "назначают вместе с другими антагонистами МКР (спиронолактон, эплеренон) и калийсберегающими "
                   "диуретиками (амилорид, триамтерен). С препаратами калия и триметопримом, в том числе в составе "
                   "ко-тримоксазола, нужны осторожность и контроль калия. По инструкции Kerendia (ЕС, раздел 4.5) "
                   "при умеренных и слабых ингибиторах CYP3A4 (эритромицин, верапамил, флувоксамин) контролируют "
                   "калий; сильные и умеренные индукторы (рифампицин, карбамазепин, зверобой, эфавиренз) снижают "
                   "концентрацию финеренона и не рекомендуются; грейпфрут и его сок исключают на время лечения. "
                   "Поэтому перед назначением нужен полный список лекарств, добавок и безрецептурных средств. "
                   "Необходимость сочетания, дозы и сроки анализов определяет врач.")
    add_para_after(S(33), "Взаимодействия. Сильные ингибиторы", "Дальше безопасность и особые группы пациентов.")
    add_para_after(S(33), "Гипокалиемия реже на финереноне",
                   "Если врач спросит о сочетаниях с другими лекарствами. Ответ: «Сильные ингибиторы CYP3A4 "
                   "противопоказаны, другие антагонисты МКР и калийсберегающие диуретики не назначают, с препаратами "
                   "калия и триметопримом контролируют калий». Это указано в инструкции Финеренон-Орвилле. "
                   "Индукторы, грейпфрут, умеренные и слабые ингибиторы описаны в инструкции Kerendia (раздел 4.5). "
                   "Конкретное сочетание оценивает врач.")
    # 42 (self-check) -> 44 (scientific section): the summary slide between them goes
    L.rep(S(42), "Итоги собираем на следующем слайде.", "На этом промоционная часть заканчивается.", notes=True)
    L.rep(S(44), "Промоционная часть закончена. Научный раздел нужен", "Научный раздел нужен", notes=True)


def remove(prs):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for n in sorted(REMOVE, reverse=True):
        el = ids[n - 1]
        prs.part.drop_rel(el.rId)
        lst.remove(el)


def main(src, dst):
    prs = Presentation(src)
    assert len(prs.slides) == 49
    edit(prs)
    remove(prs)
    assert len(prs.slides) == 46
    prs.save(dst)


if __name__ == "__main__":
    main(*sys.argv[1:])
