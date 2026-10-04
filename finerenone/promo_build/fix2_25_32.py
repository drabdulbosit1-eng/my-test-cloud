"""fix2 для слайдов 25-32 (нумерация base47). Применяется после impl_*, до перестановки и impl_final.
Правки по результатам проверки promo/check_25_32.md. Слайды адресованы по исходным номерам."""
from pptx.oxml.ns import qn
from pptx.util import Emu

import lib as L0
import r14lib as L

NBSP = " "


def _recolor(shape, mapping):
    """Replace srgbClr values inside spPr (fill, gradient stops, outline); style references stay."""
    sp = shape._element.find(qn("p:spPr"))
    n = 0
    for c in sp.iter(qn("a:srgbClr")):
        v = c.get("val").upper()
        if v in mapping:
            c.set("val", mapping[v])
            n += 1
    return n


def _fill(shape, hex_):
    sp = shape._element.find(qn("p:spPr"))
    for c in sp.find(qn("a:solidFill")).iter(qn("a:srgbClr")):
        c.set("val", hex_)


def s25(prs):
    """ADA (новый 29): замечаний к содержанию нет."""


def s26(prs):
    """Вместе с иНГЛТ2 (новый 27): заголовок переносится по смыслу, а не оставляет одно слово во 2-й строке."""
    s = L.S(prs, 26)
    L.rep(s, "чем по отдельности", "чем" + NBSP + "по" + NBSP + "отдельности", count=1)


def s27(prs):
    """Финеренон и спиронолактон (новый 28)."""
    s = L.S(prs, 27)
    # заголовок: «диабете 2 типа» переносится целиком (во 2-й строке оставалось «2 типа»)
    L0.replace_in_runs(L.title_shape(s), "диабете 2" + NBSP + "типа", "диабете" + NBSP + "2" + NBSP + "типа")
    # единое правило: у каждого процента исхода слово «комбинированный»
    L.rep(s, ", сердечно-сосудистого ниже на", ", комбинированного сердечно-сосудистого на", count=1)
    # строка исходов с запасом по высоте на случай переноса в 5 строк: +18 pt, ряды ниже сдвинуты на 18 pt
    d = 18
    for nm in ("Название_строки_3", "Финеренон_строка_3", "Спиронолактон_строка_3"):
        sh = L.shape_named(s, nm)
        sh.height = Emu(int(sh.height + d * 12700))
    for nm in ("Название_строки_4", "Финеренон_строка_4", "Спиронолактон_строка_4", "Подпись_ARTS", "Нижняя_строка"):
        sh = L.shape_named(s, nm)
        sh.top = Emu(int(sh.top + d * 12700))


# фиолетовый закреплён за финереноном: ядро, ДНК и насос получают нейтральные серые тона
GREY_MAP = {
    # ядро (градиент и контур), ядерная оболочка
    "F5EEF9": "F5F5F5", "E3D4F0": "E7E6E6", "D8C3EB": "D9D9D9", "7851AD": "7F7F7F", "C9ACDF": "BFBFBF",
    # две нити ДНК и пары оснований
    "9978BC": "8C8C8C", "61408C": "595959", "B49ACF": "BFBFBF",
    # натрий-калиевый насос (градиент и контур)
    "E7DEEE": "EDEDED", "BEAFCF": "C8C8C8", "9885AC": "8C8C8C", "847294": "6E6E6E",
}


def s28(prs):
    """Почему повышается калий (новый 33): на схеме клетки фиолетовый только у финеренона."""
    s = L.S(prs, 28)
    done = 0
    for sh in L0.iter_shapes(s.shapes):
        nm = sh.name
        if sh.has_text_frame and sh.text_frame.text.strip():
            continue
        if nm in ("Ядро", "Ядерная оболочка", "Цепь ДНК", "Пары оснований ДНК",
                  "Натрий-калиевый насос в базолатеральной мембране"):
            done += _recolor(sh, GREY_MAP)
    assert done >= 220, done


def s29(prs):
    """CYP3A4 (новый 34): замечаний нет."""


def s30(prs):
    """Взаимодействия (новый 35): в заметках фраза от третьего лица заменена обращением к МП."""
    s = L.S(prs, 30)
    L.rep(s, "Представитель объясняет, почему взаимодействия имеют значение.",
          "Объясните врачу, почему взаимодействия имеют значение.", notes=True, count=1)


def s31(prs):
    """Начало лечения (новый 30): замечаний нет."""


def s32(prs):
    """Когда контролировать калий и рСКФ (новый 31): плашки этапов нейтральные (#156082), не фиолетовые."""
    s = L.S(prs, 32)
    for nm in ("Этап_180", "Этап_480", "Этап_780"):
        _fill(L.shape_named(s, nm), "156082")
    # текст нижней карточки на одной вертикали с текстом карточек выше (x 54)
    t = L.shape_named(s, "Внеплановый_контроль")
    t.left = Emu(int(54 * 12700))
    t.width = Emu(int(848 * 12700))


def run(prs):
    for f in (s25, s26, s27, s28, s29, s30, s31, s32):
        f(prs)
