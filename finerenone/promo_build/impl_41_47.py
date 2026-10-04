"""Promo rework, slides 41-47 of base47.pptx (old numbering), source: promo/final_41_47.md.
Build: cd W && PYTHONPATH=W:W/promo python3 apply14.py base47.pptx promo/t_41.pptx impl_41_47
One function per slide: s41(prs) ... s47(prs); run(prs) calls them in order.

Order of the scientific section chosen by the customer (done in the last step, not here):
  41 -> 42 (divider) -> 47 (FINEARTS-HF) -> 45 (FINE-ONE) -> 46 (FDA, type 1 diabetes) -> 44 (FIND-CKD) -> 43 (ARTS).
Notes links are written for this order; there are no slide numbers inside slide text or notes.
Decisions of the customer that override final_41_47.md:
 - no 'Recommendations' row on the FINEARTS-HF slide (the notes only carry the developer's report on JCS/JHFS 2025);
 - FIND-CKD shares 13.9% / 16.9% only in the notes, marked 'сверить с таблицей публикации';
 - the last slide of the section (ARTS) must not promise a 'next slide' (an appendix slide follows the course).
Deviations from final_41_47.md (small, explained in impl_41_47.md):
 - card shapes start 3 pt below their plate (as on the accepted 'plate + card' slides) instead of touching it;
 - abbreviations stay at 10 pt (brief), sources 8 pt, both black; sources split into one language per line;
 - the indication on slide 41 repeats the wording of the 'start of treatment' slide (verbatim translation of the label);
 - notes of FIND-CKD: 'в максимально переносимой дозе' removed (as on the slide), first sentence of notes
   paragraph 3 does not repeat 13.9% / 16.9% (they stand once, with the remark to check)."""
import copy
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import measure as M
import r14lib as L
from lib import (add_box, add_text, bar, card, chevron, delete_shape, get_rpr, iter_shapes,
                 style_rpr, style_title)

BLACK = "000000"
PURPLE = "7030A0"   # finerenone and titles only
GREY = "7F7F7F"     # placebo
CYAN = "0F9ED5"     # comparator
TEAL = "156082"     # neutral plate
BLUE = "2C71A4"     # neutral plate
ORANGE = "E97132"   # warning / rule for the representative
CARD = "F2F2F2"
CHEV = "D9D9D9"
BODY_BOTTOM = 538   # bottom edge of the sources text in the source deck
GAP = 3             # gap between a plate and its card
NB = "\u00a0"


# ----------------------------------------------------------------------------- helpers
def nb(text):
    """Non-breaking spaces between a number and its unit, and inside 'рСКФ ≥25', '2 типа', '13 026'."""
    t = re.sub(r"(\d) (типа|мг|мл|мм|мкмоль|недел|месяц|года|лет|пациент|дн)", r"\1" + NB + r"\2", text)
    t = re.sub(r"(рСКФ|АКО|калии|калий|ФВ ЛЖ) ([≥≤<>])", r"\1" + NB + r"\2", t)
    t = re.sub(r"(\d) (\d{3})(?!\d)", r"\1" + NB + r"\2", t)
    t = re.sub(r"(1,73) (м²)", r"\1" + NB + r"\2", t)
    t = re.sub(r"(мл/мин/1,73) ", r"\1" + NB, t)
    t = re.sub(r"(\d) (ммоль)", r"\1" + NB + r"\2", t)
    t = re.sub(r"(\d) (против|и) (\d)", r"\1" + NB + r"\2 \3", t)
    return t


def geom(sh, x=None, y=None, w=None, h=None):
    if x is not None:
        sh.left = Pt(x)
    if y is not None:
        sh.top = Pt(y)
    if w is not None:
        sh.width = Pt(w)
    if h is not None:
        sh.height = Pt(h)
    return sh


def one(slide, name):
    r = [sh for sh in iter_shapes(slide.shapes) if sh.name == name]
    if len(r) != 1:
        raise KeyError(f"{name!r}: {len(r)} shapes")
    return r[0]


def drop(slide, *names):
    for n in names:
        delete_shape(one(slide, n))


def recolor_fill(sh, color):
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)


def recolor_text(sh, color=None, bold=None, size=None):
    for r in sh._element.iter(qn("a:r")):
        style_rpr(get_rpr(r), color=color, bold=bold, size=size)


def runs(text, size=15, color=BLACK, bold=False):
    """Mini markup: **bold**, ^^bold purple^^. Returns runs (text, size, bold, color)."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*|\^\^.+?\^\^)", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((nb(part[2:-2]), size, True, color))
        elif part.startswith("^^"):
            out.append((nb(part[2:-2]), size, True, PURPLE))
        else:
            out.append((nb(part), size, bold, color))
    return out


def para(text, size=15, color=BLACK, bold=False, after=None, align="l", before=None):
    d = {"runs": runs(text, size, color, bold), "align": align}
    if after is not None:
        d["space_after"] = after
    if before is not None:
        d["space_before"] = before
    return d


def need_h(paras, w, spacing=0.95):
    """Height (pt) needed by paragraphs at Arial Narrow metrics (line = size * 1.15 * spacing)."""
    tot = 0.0
    for p in paras:
        size = max(r[1] for r in p["runs"])
        n = M.n_lines([(r[0], r[2]) for r in p["runs"]], w, size)
        tot += n * size * 1.15 * spacing + (p.get("space_after") or 0) + (p.get("space_before") or 0)
    return tot


def check(label, paras, w, h, spacing=0.95, pad=0):
    need = need_h(paras, w, spacing) + pad
    flag = "OK " if need <= h + 0.5 else "!! "
    print(f"  fit {flag}{label}: need {need:.0f} of {h:.0f} pt")
    return need


def txt(slide, x, y, w, h, paras, name, anchor="t", spacing=0.95):
    check(name, paras, w, h, spacing)
    return add_text(slide, x, y, w, h, paras, anchor=anchor, name=name, line_spacing=spacing)


def plate(slide, x, y, w, h, color, text, name, size=16, bold=True, align="l"):
    return bar(slide, x, y, w, h, color, text, size=size, name=name, bold=bold, align=align)


def grey_card(slide, x, y, w, h, paras, name, anchor="t", spacing=0.95, pad_x=12, pad_y=5, fill=CARD):
    """Grey card with its text inside the shape (text cannot leave the card)."""
    check(name, paras, w - 2 * pad_x, h - 2 * pad_y, spacing)
    sh = card(slide, x, y, w, h, name=name, fill=fill)
    add_text(slide, x, y, w, h, paras, anchor=anchor, name=name, shape=sh, line_spacing=spacing)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(pad_x)
    tf.margin_top = tf.margin_bottom = Pt(pad_y)
    return sh


def block(slide, x, y, w, plate_h, card_h, color, head, paras, name, size=16, anchor="t", spacing=0.95):
    """Plate + grey card below it (gap GAP). Returns the bottom edge."""
    plate(slide, x, y, w, plate_h, color, head, f"Плашка_{name}", size=size)
    cy = y + plate_h + GAP
    grey_card(slide, x, cy, w, card_h, paras, f"Карточка_{name}", anchor=anchor, spacing=spacing)
    return cy + card_h


def section_tag(slide):
    """Label of the scientific section, right of the title (708, 24; 216x40)."""
    sh = add_box(slide, 708, 24, 216, 40, fill=TEAL, radius=8, name="Метка_раздела")
    add_text(slide, 708, 24, 216, 40,
             [{"runs": [("Научные данные:", 14, False, "FFFFFF")], "align": "c"},
              {"runs": [("другие группы пациентов", 14, False, "FFFFFF")], "align": "c"}],
             anchor="m", name="Метка_раздела", shape=sh, line_spacing=0.95)
    sh.text_frame.margin_left = sh.text_frame.margin_right = Pt(4)
    sh.text_frame.margin_top = sh.text_frame.margin_bottom = Pt(2)
    return sh


def title(slide, text, w=888, h=54):
    """Title: 20 pt purple Arial Narrow, (36, 20), at most two lines."""
    sh = L.title_shape(slide)
    L.set_lines(sh, [text])
    style_title(sh, x=36, y=20, w=w, h=h)
    n = M.n_lines([(text, False)], w, 20)
    print(f"  title: {n} line(s) at {w} pt")
    if n > 2:
        raise ValueError(f"title too long ({n} lines): {text}")
    return sh


def body_box(sh, anchor="t"):
    bp = sh.text_frame._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")
    bp.set("wrap", "square")
    bp.set("anchor", anchor)
    for k in list(bp):
        if etree.QName(k).localname in ("spAutoFit", "normAutofit", "noAutofit"):
            bp.remove(k)
    etree.SubElement(bp, qn("a:noAutofit"))


def footers(slide, abbr, sources, abbr_size=10):
    """Sources: 8 pt black, one language per paragraph, bottom edge fixed at BODY_BOTTOM, grows upwards.
    Abbreviations: 10 pt black, above the sources (at most at y=494). Returns the y above which the body must end."""
    src = L.shape_named(slide, "Источники")
    L.set_lines(src, sources)
    n = sum(M.n_lines([(p, False)], 850, 8) for p in sources)
    h = max(16.0, round(n * 9.8 + 2, 1))
    geom(src, x=36, y=BODY_BOTTOM - h, w=850, h=h)
    body_box(src, "b")
    try:
        ab = L.shape_named(slide, "Сокращения")
        L.set_lines(ab, [abbr])
    except KeyError:
        ab = add_text(slide, 36, 494, 850, 14, [{"runs": [(abbr, abbr_size, False, BLACK)], "line_spacing": 1.0}],
                      anchor="t", name="Сокращения")
    na = M.n_lines([(abbr, False)], 850, abbr_size)
    ha = round(na * 12.4, 1)
    top = min(494.0, BODY_BOTTOM - h - 3 - ha)
    geom(ab, x=36, y=top, w=850, h=ha)
    body_box(ab, "t")
    for sh in (src, ab):
        for r in sh._element.iter(qn("a:r")):
            style_rpr(get_rpr(r), color=BLACK)
    for r in ab._element.iter(qn("a:r")):
        style_rpr(get_rpr(r), size=abbr_size)
    print(f"  footers: body must end above y={top:.0f} (sources {n} line(s), abbreviations {na})")
    return top


def logo(slide, name="Логотип Орвилле"):
    geom(one(slide, name), x=898, y=500, w=56, h=34)



def edit_run(sh, old, new):
    """Replace `old` inside a single run of the shape (keeps that run's format)."""
    for r in sh._element.iter(qn("a:r")):
        t = r.find(qn("a:t"))
        txt_ = (t.text or "").replace(NB, " ")
        if old in txt_:
            t.text = nb(txt_.replace(old, new))
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            return sh
    raise KeyError(old)


def to_back(slide, sh):
    """Move a shape to the bottom of the z-order (behind the other shapes of the slide)."""
    tree = slide.shapes._spTree
    el = sh._element
    tree.remove(el)
    tree.insert(2, el)
    return sh


def set_chart_cats(chart, cats, values, fmt):
    from pptx.chart.data import CategoryChartData
    cd = CategoryChartData(number_format=fmt)
    cd.categories = cats
    cd.add_series(chart.plots[0].series[0].name, values)
    chart.replace_data(cd)

# ---- notes ------------------------------------------------------------------------------------
def notes_paras(slide):
    return [p.text for p in slide.notes_slide.notes_text_frame.paragraphs if p.text.strip()]


def set_notes_paras(slide, paras):
    L.set_notes(slide, "\n\n".join(paras))


def ask(q, a):
    return f"Если врач спросит: «{q}» Ответ: {a}"


def phrase(text):
    return f"Фраза для визита: «{text}»"


# ---- charts -----------------------------------------------------------------------------------
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _c(tag):
    return "{%s}%s" % (C_NS, tag)


def chart_of(slide, name):
    sh = one(slide, name)
    if not sh.has_chart:
        raise KeyError(f"{name!r} is not a chart")
    return sh, sh.chart


def label_text_element(idx, text, size, color=None, bold=False):
    """c:dLbl with a fixed text (so decimal comma and minus sign do not depend on regional settings)."""
    fill = f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>' if color else ""
    xml = (f'<c:dLbl xmlns:c="{C_NS}" xmlns:a="{A_NS}"><c:idx val="{idx}"/><c:tx><c:rich><a:bodyPr/><a:lstStyle/>'
           f'<a:p><a:pPr><a:defRPr sz="{int(size * 100)}" b="{int(bold)}"/></a:pPr>'
           f'<a:r><a:rPr lang="ru-RU" sz="{int(size * 100)}" b="{int(bold)}">{fill}'
           f'<a:latin typeface="Arial Narrow"/><a:cs typeface="Arial Narrow"/></a:rPr><a:t>{text}</a:t></a:r></a:p>'
           f'</c:rich></c:tx><c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
           f'<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>')
    return etree.fromstring(xml)


# ----------------------------------------------------------------------------- slide 41
def s41(prs):
    s = L.S(prs, 41)
    title(s, "Финеренон снижает риск почечных и сердечно-сосудистых исходов у взрослых с ХБП и диабетом 2 типа; "
             "контроль калия сводится к простому правилу")

    for n in ("Итог_полоса_1", "Итог_полоса_2", "Итог_полоса_3", "Итог_текст_1", "Итог_текст_2", "Итог_текст_3"):
        drop(s, n)

    y_pl, h_pl, h_card = 84, 30, 260
    cols = [(36, TEAL, "Кого искать"), (336, BLUE, "Что доказано"), (636, ORANGE, "Как назначать и контролировать")]
    sz, aft = 15, 6
    c1 = [para("**Взрослые с ХБП, связанной с диабетом 2 типа.**", sz, after=aft),
          para("**Показание:** снижение риска устойчивого снижения скорости клубочковой фильтрации, терминальной "
               "стадии болезни почек, смерти от осложнений сердечно-сосудистых заболеваний, нефатального инфаркта "
               "миокарда и госпитализации по поводу СН.", sz, after=aft),
          para("**В исследованиях:** альбуминурия (АКО от 30 мг/г) на ингибиторе АПФ или БРА в максимально "
               "переносимой дозе.", sz)]
    c2 = [para("По сравнению с плацебо риск комбинированного почечного исхода ниже на ^^18%^^ (17,8% против 21,1%; "
               "FIDELIO-DKD, медиана 2,6 года), комбинированного сердечно-сосудистого исхода ниже на ^^13%^^ "
               "(12,4% против 14,2%; FIGARO-DKD, медиана 3,4 года).", sz, after=aft),
          para("**Госпитализации по поводу СН в FIGARO-DKD:** риск ниже на 29% (ОР 0,71).", sz, after=aft),
          para("**FIDELITY, 13 026 пациентов:** сердечно-сосудистый исход ОР 0,86, почечный исход ОР 0,77.", sz)]
    c3 = [para("**Старт:** калий ≤4,8 ммоль/л (4,8–5,0 с дополнительным контролем), выше 5,0 не начинают; "
               "рСКФ ≥25. Доза 10 или 20 мг по рСКФ, цель 20 мг один раз в сутки.", sz, after=aft),
          para("Калий и рСКФ до начала, через 4 недели после начала, возобновления или повышения дозы, "
               "далее периодически.", sz, after=aft),
          para("**Окончательная отмена из-за гиперкалиемии:** 1,7% против 0,6% на плацебо (FIDELITY).", sz, after=aft),
          para("До назначения проверяют противопоказания и взаимодействия.", sz)]
    for (x, color, head), paras, nm in zip(cols, (c1, c2, c3), ("Кого искать", "Что доказано", "Как назначать")):
        block(s, x, y_pl, 288, h_pl, h_card - GAP, color, head, paras, nm, size=18)

    # bottom strip: the question for the physician
    plate(s, 36, 388, 888, 26, TEAL, "Вопрос врачу на визите", "Плашка_Вопрос врачу", size=16)
    grey_card(s, 36, 388 + 26 + GAP, 888, 58 - GAP, [para("У каких ваших пациентов с ХБП и диабетом 2 типа "
              "альбуминурия сохраняется, несмотря на ингибитор АПФ или БРА?", 18)], "Карточка_Вопрос врачу",
              anchor="m")

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; АКО: отношение "
            "альбумина к креатинину мочи; СН: сердечная недостаточность; АПФ: ангиотензинпревращающий фермент; "
            "БРА: блокатор рецепторов ангиотензина II; ОР: отношение рисков.",
            ["Bakris GL, et al. N Engl J Med. 2020;383:2219–2229. Pitt B, et al. N Engl J Med. 2021;385:2252–2263. "
             "Agarwal R, et al. Eur Heart J. 2022;43:474–484.",
             "Финеренон-Орвилле. Инструкция по медицинскому применению."])

    # ---- notes
    L.rep(s, "По предоставленной местной инструкции Финеренон-Орвилле предназначен для взрослых с ХБП, связанной с "
             "диабетом 2 типа, для снижения перечисленных в показании почечных и сердечно-сосудистых рисков. "
             "Исследовательскую группу и критерии отбора, включая альбуминурию, нужно описывать отдельно от текста "
             "показания.",
          "Вернёмся к главному из занятия. Финеренон-Орвилле показан взрослым с ХБП, связанной с диабетом 2 типа: он "
          "снижает риск устойчивого снижения скорости клубочковой фильтрации, терминальной стадии болезни почек, "
          "смерти от осложнений сердечно-сосудистых заболеваний, нефатального инфаркта миокарда и госпитализации "
          "по поводу сердечной недостаточности. В исследованиях участвовали пациенты с альбуминурией на ингибиторе "
          "АПФ или БРА. Группу исследования и текст показания называйте раздельно.", notes=True)
    L.rep(s, "Эффекты относятся к соответствующим группам и определениям исходов. Снижение АКО не позволяет "
             "рассчитать пользу для отдельного пациента.",
          "В объединённом анализе FIDELITY (13 026 пациентов) риск сердечно-сосудистого исхода ниже на 14% "
          "(ОР 0,86; ЧБНЛ 46 за 3 года), риск почечного исхода ниже на 23% (ОР 0,77; ЧБНЛ 60). ЧБНЛ: число "
          "больных, которых нужно лечить, чтобы предотвратить одно событие.", notes=True)
    L.rep(s, "Блокада минералокортикоидного рецептора может повышать калий. До начала лечения проверяют "
             "противопоказания, взаимодействия, калий и рСКФ. Затем анализы повторяют через четыре недели после "
             "начала, возобновления или повышения дозы и далее периодически. Решение о лечении принимает врач. "
             "Медицинскому отделу передают вопросы о применении у конкретного пациента, в том числе вне местного "
             "показания, а в фармаконадзор сообщают о нежелательных явлениях.",
          "Правило по калию простое: анализ до начала, через четыре недели после начала, возобновления или "
          "повышения дозы, далее периодически. Начинают при калии до 4,8 ммоль/л; при 4,8–5,0 можно начинать с "
          "дополнительным контролем в первые четыре недели; выше 5,0 не начинают. В FIDELITY гиперкалиемия была у "
          "14,0% на финереноне и у 6,9% на плацебо; окончательная отмена из-за неё у 1,7% и 0,6%. До назначения "
          "проверяют противопоказания и взаимодействия. Решение о лечении принимает врач. О нежелательных явлениях "
          "сообщают в фармаконадзор. Вопросы о применении вне показания передают в медицинский отдел.", notes=True)
    story = notes_paras(s)
    if len(story) != 3:
        raise KeyError(f"s41 notes: expected 3 paragraphs, got {len(story)}")
    set_notes_paras(s, story + [
        "Дальше научный раздел: данные по другим группам пациентов. Сами эти темы на визите вы не начинаете.",
        phrase("У каких ваших пациентов с ХБП и диабетом 2 типа альбуминурия сохраняется на ингибиторе АПФ или БРА? "
               "У таких пациентов финеренон снижал риск почечных и сердечно-сосудистых исходов по сравнению с "
               "плацебо. Калий проверяем до начала и через 4 недели."),
        ask("Какая польза у моего пациента?", "цифры описывают группы в исследованиях, решение о лечении "
            "конкретного пациента принимает врач."),
        ask("Это относительное снижение?", "да, относительное; абсолютные доли даны рядом (17,8% и 21,1%; 12,4% и "
            "14,2%)."),
    ])


# ----------------------------------------------------------------------------- slide 42
def s42(prs):
    s = L.S(prs, 42)
    # title = name of the section (20 pt, one line); subtitle 16 pt
    title(s, "Научные данные: другие группы пациентов", w=888, h=36)
    for n in ("Связь ARTS с сердцем", "ARTS точка", "Связь FINEARTS-HF с сердцем", "FINEARTS-HF точка",
              "Связь FIND-CKD с почкой", "FIND-CKD точка", "Связь FINE-ONE с почкой", "FINE-ONE точка",
              "ARTS", "ARTS описание", "FINEARTS-HF", "FINEARTS-HF описание", "FIND-CKD", "FIND-CKD описание",
              "FINE-ONE", "FINE-ONE описание", "Оговорка о местной инструкции"):
        drop(s, n)
    txt(s, 36, 58, 888, 38,
        [para("Группы пациентов вне показания Финеренон-Орвилле. Данные нужны, чтобы понять вопрос врача и передать "
              "его точно. Материал для обучения МП: врачам не показывать и не передавать.", 16)],
        "Подзаголовок", spacing=0.95)

    # rule for the representative: orange plate, grey card, three white steps joined by chevrons
    plate(s, 36, 100, 888, 28, ORANGE, "Правило для медицинского представителя", "Плашка_Правило для МП", size=16)
    grey_card(s, 36, 128 + GAP, 888, 100 - GAP, [], "Карточка_Правило для МП")
    txt(s, 50, 136, 860, 24, [para("**МП эти темы сам не начинает.**", 18)], "Правило_строка", anchor="m")
    steps = ["Врач спрашивает о другой группе пациентов",
             "МП не обсуждает применение и записывает вопрос",
             "Медицинский отдел или медицинский советник отвечает врачу"]
    for i, (x, t) in enumerate(zip((50, 342, 634), steps), 1):
        grey_card(s, x, 166, 266, 56, [para(t, 15, align="c")], f"Шаг_{i}", anchor="m", fill="FFFFFF", pad_x=8, pad_y=3)
    for i, x in enumerate((318, 610), 1):
        chevron(s, x, 179, w=22, h=30, name=f"Шеврон_{i}", color=CHEV)

    # map of the section: four 'plate + card' blocks, picture of heart and kidney between the columns
    items = [
        (36, 238, "FINEARTS-HF", "Сердечная недостаточность с ФВ ЛЖ ≥40%, 6001 пациент. Зарегистрировано, включая США, "
                                 "Евросоюз, Японию, Великобританию."),
        (540, 238, "FINE-ONE", "ХБП при диабете 1 типа, 242 пациента. Зарегистрировано в США (16.09.2026)."),
        (36, 358, "FIND-CKD", "ХБП без диабета, 1584 пациента. Нигде в мире не зарегистрировано."),
        (540, 358, "ARTS", "Сердечная недостаточность со сниженной ФВ ЛЖ и ХБП, II фаза, 4 недели: повышение "
                           "калия, сравнение со спиронолактоном."),
    ]
    for x, y, head, text in items:
        block(s, x, y, 384, 26, 84 - GAP, TEAL, head, [para(text, 15)], head, size=16, anchor="m")
    pic = one(s, "Рисунок 24")
    geom(pic, x=432, y=238, w=96, h=230)

    footers(s,
            "МП: медицинский представитель; СН: сердечная недостаточность; ХБП: хроническая болезнь почек; "
            "ФВ ЛЖ: фракция выброса левого желудочка.",
            ["Финеренон-Орвилле. Инструкция по медицинскому применению. Статусы регистрации на 03.10.2026: решения "
             "регуляторов и сообщения разработчика; источники по каждому исследованию на следующих слайдах."])
    logo(s, "Рисунок 45")

    # ---- notes
    L.rep(s, "Рассмотрим исследования в других клинических ситуациях: короткое сравнение со спиронолактоном, ХБП без "
             "диабета, ХБП при диабете 1 типа и сердечная недостаточность с фракцией выброса 40% и выше. У них "
             "другие группы пациентов, задачи и иногда дозы.",
          "Промоционная часть закончена. Дальше научный раздел: данные по другим группам пациентов. Он нужен, чтобы "
          "вы понимали вопрос врача и передавали его точно. Здесь сердечная недостаточность с ФВ ЛЖ 40% и выше, "
          "ХБП при диабете 1 типа, ХБП без диабета и короткое сравнение со спиронолактоном при сердечной "
          "недостаточности. У этих исследований другие пациенты, цели и иногда дозы.", notes=True)
    L.rep(s, "Публикация результата и зарубежное регистрационное решение не меняют инструкцию Финеренон-Орвилле в "
             "Узбекистане. В предоставленном местном тексте группа пациентов ограничена взрослыми с ХБП, связанной с "
             "диабетом 2 типа. При этом снижение риска госпитализации по поводу сердечной недостаточности входит в "
             "местное показание для этой группы. Оно отличается от самостоятельного показания при сердечной "
             "недостаточности независимо от ХБП при диабете 2 типа.",
          "Правило раздела: эти темы вы сами на визите не начинаете. Если врач спросил, вы не обсуждаете применение: "
          "записываете вопрос и передаёте его в медицинский отдел или медицинскому советнику. Что сказать врачу: "
          "«Спасибо за вопрос. Такие темы разбирает медицинский отдел. Я передам вопрос, и медицинский советник "
          "свяжется с вами». Материалы раздела врачу не показывают и не отправляют. Показание Финеренон-Орвилле: "
          "взрослые с ХБП, связанной с диабетом 2 типа. Снижение риска госпитализации по поводу сердечной "
          "недостаточности входит в него для этой группы. Сердечная недостаточность с ФВ ЛЖ 40% и выше без ХБП при "
          "диабете 2 типа в показание не входит.", notes=True)
    L.rep(s, "Эти сведения помогают точно понять вопрос врача и передать запрос медицинскому отделу. Учебные "
             "материалы используют и передают по действующей процедуре компании.",
          "Статусы регистрации даны на 03.10.2026, их обновляет медицинский отдел. Когда показание появится в "
          "инструкции Финеренон-Орвилле, соответствующий слайд переносится в промоционную часть. Учебные материалы "
          "используют и передают по действующей процедуре компании. Начнём с самой крупной группы: сердечная "
          "недостаточность с ФВ ЛЖ 40% и выше.", notes=True)


# ----------------------------------------------------------------------------- slide 43
def s43(prs):
    s = L.S(prs, 43)
    title(s, "ARTS, II фаза: при СН со сниженной ФВ ЛЖ и ХБП калий за 4 недели повышался реже на финереноне (5,3%), "
             "чем на спиронолактоне (12,7%)", w=660)
    section_tag(s)

    # chart: placebo, finerenone, spironolactone (ascending), colours by meaning, decimal comma
    sh, ch = chart_of(s, "Chart 4")
    set_chart_cats(ch, ["Плацебо", "Финеренон", "Спиронолактон"], (0.015, 0.053, 0.127), "0.0%")
    ser = ch._chartSpace.find(".//" + _c("ser"))
    dpts = ser.findall(_c("dPt"))
    dlbls = ser.find(_c("dLbls"))
    lbls = dlbls.findall(_c("dLbl"))
    if len(dpts) != 2 or len(lbls) != 2:
        raise KeyError("s43 chart: unexpected dPt/dLbl structure")

    def recolor(el, color):
        for c in el.iter("{%s}srgbClr" % A_NS):
            c.set("val", color)

    def reindex(el, idx):
        el.find(_c("idx")).set("val", str(idx))
        ext = el.find(_c("extLst"))
        if ext is not None:
            el.remove(ext)

    p_f, p_s = dpts           # finerenone (purple), spironolactone (cyan)
    l_f, l_s = lbls
    p_p, l_p = copy.deepcopy(p_f), copy.deepcopy(l_f)
    recolor(p_p, GREY)
    recolor(l_p, GREY)
    reindex(p_p, 0)
    reindex(l_p, 0)
    reindex(p_f, 1)
    reindex(l_f, 1)
    reindex(p_s, 2)
    reindex(l_s, 2)
    p_f.addprevious(p_p)
    l_f.addprevious(l_p)
    geom(sh, x=45, y=140, w=505, h=186)

    # caption under the chart
    txt(s, 42, 328, 510, 36,
        [para("p=0,048 (финеренон и спиронолактон); p=0,32 (финеренон и плацебо). Группы доз финеренона объединены "
              "(дополнительный анализ).", 14)], "Подпись диаграммы", spacing=0.95)

    # conclusion card on the left: placebo moved into the chart, mean potassium rise added
    L.rep(s, "Плацебо в ARTS: 1,5%. При лечении обоими препаратами контролируют калий и функцию почек.",
          nb("Средний прирост калия за 4 недели: 0,04–0,30 ммоль/л на финереноне и 0,45 ммоль/л на спиронолактоне. "
             "Калий и функцию почек контролируют при обоих препаратах."))
    card_l = one(s, "Карточка_вывод")
    txt_l = L.shape_with(s, "В ARTS исследователи")
    geom(card_l, y=368, h=102)
    geom(txt_l, x=48, y=374, w=516, h=90)
    recolor_text(txt_l, size=15)
    txt_l.text_frame.paragraphs[0].space_after = Pt(4)
    check("Вывод", [para("В ARTS исследователи реже сообщали о повышении калия при финереноне, чем при "
                         "спиронолактоне.", 15, after=4),
                    para("Средний прирост калия за 4 недели: 0,04–0,30 ммоль/л на финереноне и 0,45 ммоль/л на "
                         "спиронолактоне. Калий и функцию почек контролируют при обоих препаратах.", 15)], 516, 90)

    # right column
    L.rep(s, "392 пациента с сердечной недостаточностью со сниженной фракцией выброса и умеренным нарушением функции "
             "почек (анализ безопасности).",
          nb("392 пациента с сердечной недостаточностью, ФВ ЛЖ 40% и ниже, рСКФ от 30 до 60 мл/мин/1,73 м² "
             "(часть B, анализ безопасности)."))
    L.rep(s, "Дополнительный анализ объединённых групп доз. Спиронолактон применяли открыто. Изучены указанные дозы "
             "и 4 недели лечения. Долгосрочную сравнительную безопасность не оценивали.",
          "Группы доз финеренона объединены (дополнительный анализ). Спиронолактон применяли открыто.")
    who = L.shape_with(s, "Кто участвовал")
    recolor_text(who, size=16)
    geom(who, y=126, h=100)
    check("Кто участвовал", [para("**Кто участвовал:** 392 пациента с сердечной недостаточностью, ФВ ЛЖ 40% и ниже, "
                                  "рСКФ от 30 до 60 мл/мин/1,73 м² (часть B, анализ безопасности).", 16),
                             para("**Наблюдение:** лечение 4 недели.", 16)], 310, 100)
    geom(one(s, "Карточка_ARTS"), h=116)
    plate_d = one(s, "Плашка_дозы")
    geom(plate_d, y=246)
    geom(one(s, "Карточка_дозы"), y=280, h=72)
    geom(L.shape_with(s, "Финеренон: 2,5"), y=288)
    geom(one(s, "Карточка_ограничение"), y=362, h=108)
    lim = L.shape_with(s, "Спиронолактон применяли открыто")
    geom(lim, y=362, h=108)
    lim.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    drop(s, "Статус популяции")

    footers(s,
            "ХБП: хроническая болезнь почек; СН: сердечная недостаточность; ФВ ЛЖ: фракция выброса левого желудочка; "
            "рСКФ: расчётная скорость клубочковой фильтрации.",
            ["Pitt B, et al. Eur Heart J. 2013;34:2453–2463."])

    # ---- notes
    L.rep(s, "В части B ARTS участвовали пациенты",
          "Последний слайд раздела: справка к вопросу «чем финеренон отличается от спиронолактона». "
          "В части B ARTS участвовали пациенты", notes=True)
    L.rep(s, "Результат показывает различие в переносимости изученных схем за четыре недели. Клинические исходы и "
             "долгосрочное преимущество по безопасности здесь не оценивали. При обоих препаратах контролируют калий "
             "и функцию почек.",
          "Результат показывает различие в переносимости изученных схем за четыре недели. Другие данные ARTS: "
          "нарушение функции почек как нежелательное явление у 3,8% на финереноне, у 28,6% на спиронолактоне и у "
          "9,2% на плацебо; BNP, NT-proBNP и АКО снижались на финереноне не меньше, чем на спиронолактоне "
          "(описательный анализ, значимого общего эффекта лечения не было). Калий и функцию почек контролируют при "
          "любом антагонисте минералокортикоидного рецептора.", notes=True)
    story = notes_paras(s)
    if len(story) != 2:
        raise KeyError(f"s43 notes: expected 2 paragraphs, got {len(story)}")
    set_notes_paras(s, story + [
        "Раздел закончен. Вопросы врача по этим темам передавайте в медицинский отдел или медицинскому советнику.",
        ask("А долгосрочная безопасность и клинические исходы?", "в ARTS их не оценивали. Это исследование II фазы, "
            "оно подбирало дозы для программы III фазы; дозы финеренона в ARTS были от 2,5 до 10 мг."),
        ask("А гинекомастия? А эплеренон?", "ARTS эти вопросы не изучал, сравнения с эплереноном в нём не было."),
    ])


# ----------------------------------------------------------------------------- slide 44
def s44(prs):
    s = L.S(prs, 44)
    title(s, "FIND-CKD: у взрослых с ХБП без диабета рСКФ на финереноне снижалась медленнее, чем на плацебо "
             "(−3,3 против −4,0 мл/мин/1,73 м² в год)", w=660)
    section_tag(s)

    # patients and comparison: plate + card
    drop(s, "Пациенты и сравнение")
    plate(s, 36, 80, 888, 22, TEAL, "Пациенты и сравнение", "Плашка_Пациенты и сравнение", size=16)
    grey_card(s, 36, 80 + 22 + GAP, 888, 146 - 105, [para(
        "1584 взрослых с ХБП без диабета: рСКФ от 25 до менее 90 мл/мин/1,73 м², АКО от 200 до 3500 мг/г. "
        "Финеренон 10 или 20 мг либо плацебо на фоне ингибитора ренин-ангиотензиновой системы.", 15)],
        "Карточка_Пациенты и сравнение", anchor="m", pad_y=3)

    # primary outcome caption (units moved into the same line)
    drop(s, "Единицы графика")
    cap = one(s, "Годовой темп")
    L.rep(s, "до 32-го месяца", "до 32-го месяца, мл/мин/1,73 м² в год")
    recolor_text(cap, size=16)
    geom(cap, y=152, h=22)
    check("Годовой темп", [para("**Первичный исход: среднегодовое изменение рСКФ до 32-го месяца, мл/мин/1,73 м² в "
                                "год**", 16)], 888, 22)

    # chart: fixed label texts with decimal comma and minus sign
    sh, ch = chart_of(s, "Годовое изменение рСКФ")
    geom(sh, x=36, y=176, w=540, h=130)
    dl = ch._chartSpace.find(".//" + _c("ser")).find(_c("dLbls"))
    first = dl[0]
    first.addprevious(label_text_element(0, "−3,3", 16))
    first.addprevious(label_text_element(1, "−4,0", 16))
    ci = one(s, "ДИ групп")
    recolor_text(ci, size=14)
    geom(ci, x=36, y=308, w=566, h=18)

    # difference: grey card with the big number
    drop(s, "Разница годовых темпов", "ДИ годового темпа")
    grey_card(s, 600, 176, 324, 150,
              [para("^^Разница 0,7^^", 32, after=4),
               para("мл/мин/1,73 м² в год в пользу финеренона", 18, after=2),
               para("95% ДИ 0,3–1,1; p<0,001", 18)], "Карточка_разница", anchor="m")

    # secondary outcome, composition, safety: one grey card
    drop(s, "Вторичный исход", "Состав вторичного исхода", "Безопасность и статус")
    grey_card(s, 36, 330, 888, 96,
              [para("**Вторичный комбинированный исход (почечные и сердечно-сосудистые события):** риск ниже на 23% по "
                    "сравнению с плацебо; ОР 0,77 (95% ДИ 0,60–0,99), p=0,04.", 15, after=3),
               para("**Состав:** снижение рСКФ на 57% и более от исходного, почечная недостаточность, госпитализация "
                    "по поводу СН или сердечно-сосудистая смерть.", 15, after=3),
               para("**Гиперкалиемия:** 17,0% на финереноне и 13,3% на плацебо; окончательная отмена из-за неё: 1,5% "
                    "и 0,1%.", 15)], "Карточка_вторичный исход и безопасность", anchor="m", pad_y=4)

    # registration status
    plate(s, 36, 430, 888, 22, TEAL, "Статус регистрации на 03.10.2026", "Плашка_Статус регистрации", size=16)
    grey_card(s, 36, 430 + 22 + GAP, 888, 480 - 455, [para("Показание при ХБП без диабета нигде в мире не "
              "зарегистрировано.", 15)], "Карточка_Статус регистрации", anchor="m", pad_y=2)

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; АКО: отношение "
            "альбумина к креатинину мочи; СН: сердечная недостаточность; ДИ: доверительный интервал; ОР: отношение "
            "рисков.",
            ["Heerspink HJL, et al. N Engl J Med. 2026;395:533–545.",
             "Статус регистрации на 03.10.2026: решения регуляторов и сообщения Bayer AG."])

    # ---- notes
    L.rep(s, "FIND-CKD включало 1584 взрослых", "Теперь ХБП без диабета. FIND-CKD включало 1584 взрослых", notes=True)
    L.rep(s, "Финеренон сравнивали с плацебо на фоне ингибитора АПФ или блокатора рецепторов ангиотензина в "
             "максимально переносимой дозе.",
          "Финеренон 10 или 20 мг сравнивали с плацебо на фоне ингибитора ренин-ангиотензиновой системы.",
          notes=True)
    L.rep(s, "Это более медленное среднегодовое снижение фильтрации. Показатель рассчитан по модели за весь период "
             "оценки и не равен изменению рСКФ в отдельно взятой точке. Он не доказывает восстановления "
             "повреждённой ткани почки.",
          "Среднегодовое снижение фильтрации на финереноне медленнее. Показатель рассчитан по модели за весь период "
          "оценки, а не по одной точке. После отмены препарата наклон составил +1,2 на финереноне и −0,5 на плацебо "
          "(рисунок 1A публикации).", notes=True)
    L.rep(s, "Вторичный комбинированный исход возник у 13,9% пациентов на финереноне и у 16,9% на плацебо. Его "
             "проверяли по заранее заданной иерархии. Результат нельзя переносить на каждый компонент отдельно.",
          "Вторичный комбинированный исход (почечные и сердечно-сосудистые события): ОР 0,77 (95% ДИ 0,60–0,99), "
          "p=0,04; риск ниже на 23% по сравнению с плацебо. Его проверяли по заранее заданной иерархии. Доли "
          "пациентов с событием по сообщению о докладе на ERA 2026: 13,9% на финереноне и 16,9% на плацебо (сверить "
          "с таблицей публикации).", notes=True)
    L.rep(s, "В предоставленной местной инструкции показание при ХБП без диабета не указано. Результат исследования "
             "сам по себе не меняет регистрационное показание.",
          "Статус на 03.10.2026: показание при ХБП без диабета нигде в мире не зарегистрировано. По сообщениям "
          "разработчика, в августе 2026 поданы заявки на расширение применения при ХБП в Китае (ХБП без диабета) и в "
          "Японии (28.08.2026); решений нет. Рекомендаций с этим показанием на дату подготовки не найдено. Вопросы "
          "врача по ХБП без диабета передавайте в медицинский отдел.", notes=True)
    story = notes_paras(s)
    if len(story) != 3:
        raise KeyError(f"s44 notes: expected 3 paragraphs, got {len(story)}")
    set_notes_paras(s, story + [
        "Последний слайд раздела: сравнение со спиронолактоном по калию, ARTS.",
        ask("А компоненты вторичного исхода?", "почечный комбинированный исход: ОР 0,78 (95% ДИ 0,60–1,01); "
            "сердечно-сосудистый комбинированный исход: ОР 0,60 (95% ДИ 0,27–1,33). Доверительные интервалы широкие "
            "и включают 1, поэтому основу составляет общий исход."),
        ask("А безопасность?", "гиперкалиемия потребовала госпитализации у 0,9% на финереноне и у 0,6% на плацебо."),
    ])


# ----------------------------------------------------------------------------- slide 45
def s45(prs):
    s = L.S(prs, 45)
    title(s, "FINE-ONE: при ХБП и диабете 1 типа АКО за 6 месяцев снизилось на 34% на финереноне и на 12% на плацебо",
          w=660)
    section_tag(s)

    # who took part / what was compared: plate + grey card
    drop(s, "Text_4")
    plate(s, 36, 80, 888, 22, TEAL, "Пациенты и сравнение", "Плашка_Пациенты и сравнение", size=16)
    grey_card(s, 36, 80 + 22 + GAP, 888, 163 - 105,
              [para("**Кто участвовал:** 242 взрослых с хронической болезнью почек и диабетом 1 типа; рСКФ ≥25 и <90 "
                    "мл/мин/1,73 м², АКО ≥200 и <5000 мг/г.", 16, after=2),
               para("**Что сравнивали:** финеренон 10 или 20 мг (по рСКФ) либо плацебо на фоне инсулина и ингибитора "
                    "АПФ или БРА.", 16)], "Карточка_Пациенты и сравнение", anchor="m", pad_y=3)

    # chart caption and chart
    cap = one(s, "Text_5")
    L.rep(s, "Среднее по визитам на 3-м и 6-м месяцах", "За 6 месяцев (среднее по визитам на 3-м и 6-м месяцах)")
    geom(cap, x=36, y=168, w=540, h=40)
    sh, ch = chart_of(s, "Chart 6")
    geom(sh, x=38, y=208, w=479, h=170)

    # right block: big number 25% (32 pt) and explanation
    geom(one(s, "Карточка_результат"), x=580, y=168, w=344, h=210)
    big = one(s, "Text_7")
    L.rep(s, "−25%", "25%")
    recolor_text(big, size=32)
    geom(big, x=596, y=180, w=312, h=44)
    expl = one(s, "Text_8")
    L.rep(s, "Относительная разница с плацебо по первичному исходу.",
          "Первичный исход достигнут: АКО снизилось на 25% больше, чем на плацебо (относительная разница между "
          "группами).")
    geom(expl, x=596, y=230, w=312, h=142)
    expl.text_frame.paragraphs[0].space_after = Pt(6)
    check("Блок справа", [para("Первичный исход достигнут: АКО снизилось на 25% больше, чем на плацебо (относительная "
                               "разница между группами).", 18),
                          para("Отношение геометрических средних 0,75 (95% ДИ 0,65–0,87), p<0,001.", 18)], 312, 146)

    # conclusions under the chart: one grey card
    drop(s, "Text_9", "Text_10")
    grey_card(s, 36, 381, 888, 44,
              [para("В FINE-ONE оценивали АКО; клинические исходы не изучали.", 15, after=2),
               para("**Гиперкалиемия:** 10,1% на финереноне и 3,3% на плацебо; окончательная отмена из-за неё у 1,7% на "
                    "финереноне.", 15)], "Карточка_выводы", anchor="m", pad_y=3)

    # registration status
    plate(s, 36, 430, 888, 22, TEAL, "Статус регистрации на 03.10.2026", "Плашка_Статус регистрации", size=16)
    grey_card(s, 36, 430 + 22 + GAP, 888, 481 - 455,
              [para("США (FDA): 16.09.2026, подробности на следующем слайде. Решений других регуляторов по сообщениям "
                    "нет.", 15)], "Карточка_Статус регистрации", anchor="m", pad_y=2)

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; АКО: отношение "
            "альбумина к креатинину мочи; АПФ: ангиотензинпревращающий фермент; БРА: блокатор рецепторов "
            "ангиотензина II; ДИ: доверительный интервал.",
            ["Heerspink HJL, et al. N Engl J Med. 2026;394:947–957. FDA. Supplement Approval Letter, NDA 215341/S-011, "
             "16.09.2026."])

    # ---- notes
    L.rep(s, "В FINE-ONE участвовали 242 взрослых", "От сердечной недостаточности переходим к почкам: ХБП при "
          "диабете 1 типа. В FINE-ONE участвовали 242 взрослых", notes=True)
    L.rep(s, "Высоты столбиков нельзя просто вычитать, поскольку они показывают изменения внутри групп.",
          "Высоты столбиков нельзя просто вычитать, поскольку они показывают изменения внутри групп. Поэтому "
          "разницу 34% и 12% не читают как 22 процентных пункта: эффект выражен отношением геометрических средних "
          "0,75, то есть на 25% больше снижения.", notes=True)
    L.rep(s, "Оценивали промежуточный лабораторный показатель. Диализ и другие клинические исходы непосредственно в "
             "FINE-ONE не изучали. Гиперкалиемия отмечалась у 10,1% пациентов на финереноне и у 3,3% на плацебо. В "
             "предоставленном местном показании ХБП при диабете 1 типа не указана. Зарубежное регистрационное решение "
             "относится к отдельному документу и рассматривается далее.",
          "В FINE-ONE оценивали АКО. Диализ и другие клинические исходы не изучали, поэтому показание FDA "
          "сформулировано через ожидаемый эффект (следующий слайд). Гиперкалиемия: 10,1% на финереноне и 3,3% на "
          "плацебо; двое пациентов (1,7%) окончательно прекратили приём финеренона из-за неё. К 6-му месяцу рСКФ "
          "снизилась на 5,6 мл/мин/1,73 м² на финереноне и на 2,7 на плацебо (разница −2,9; 95% ДИ −5,1…−0,7); в "
          "период отмены значения приближались к исходным. Вопросы врача о применении при диабете 1 типа передавайте "
          "в медицинский отдел. Решение FDA рассматриваем на следующем слайде.", notes=True)


# ----------------------------------------------------------------------------- slide 46
def s46(prs):
    s = L.S(prs, 46)
    title(s, "16.09.2026 FDA одобрило финеренон при ХБП и диабете 1 типа на основании снижения АКО в FINE-ONE и данных "
             "при диабете 2 типа", w=660)
    section_tag(s)

    # block 'Решение FDA'
    edit_run(one(s, "США_текст_1"), "16.09.2026.", "16.09.2026 (объявлено 17.09.2026).")
    # block 'Показание в США': basis of approval shortened
    edit_run(one(s, "США_текст_2"),
             "снижение АКО (суррогатная конечная точка) в FINE-ONE и данные исследований при диабете 2 типа. "
             "Клинические исходы (устойчивое снижение рСКФ, терминальная почечная недостаточность) в FINE-ONE не "
             "оценивали.",
             "снижение АКО в FINE-ONE (суррогатная конечная точка) и данные исследований при диабете 2 типа.")
    # bottom block: status instead of 'Местная инструкция', neutral colour instead of orange
    pl3 = one(s, "США_полоса_3")
    recolor_fill(pl3, TEAL)
    L.rep(s, "Местная инструкция", "Статус регистрации на 03.10.2026")
    L.rep(s, "В предоставленной инструкции Финеренон-Орвилле показание при ХБП, связанной с диабетом 1 типа, не "
             "указано. Зарубежная регистрация не меняет местное показание.",
          "США: одобрено 16.09.2026. Решений других регуляторов по сообщениям нет; актуальный статус даёт "
          "медицинский отдел.")

    footers(s,
            "ХБП: хроническая болезнь почек; рСКФ: расчётная скорость клубочковой фильтрации; АКО: отношение "
            "альбумина к креатинину мочи; FDA: Управление по контролю качества пищевых продуктов и лекарств США.",
            ["FDA. Supplement Approval Letter, NDA 215341/S-011, 16.09.2026. Kerendia. Prescribing Information. "
             "US FDA."])

    # ---- notes
    L.rep(s, "16 сентября 2026 года FDA одобрило", "Результат FINE-ONE привёл к регистрационному решению. "
          "16 сентября 2026 года FDA одобрило", notes=True)
    L.rep(s, "Решение FDA действует в США и относится к указанному препарату. Оно не меняет предоставленную "
             "инструкцию Финеренон-Орвилле для Узбекистана, где такого показания нет. Запрос врача о применении при "
             "диабете 1 типа передают медицинскому отделу по процедуре компании.",
          "Решение FDA действует в США и относится к препарату Керендия. В инструкции Финеренон-Орвилле такого "
          "показания нет. Если врач спросит о диабете 1 типа, вы не обсуждаете применение. Что сказать: «Спасибо за "
          "вопрос. Такие темы разбирает медицинский отдел. Я передам вопрос, и медицинский советник свяжется с "
          "вами». Решение о применении принимает врач.", notes=True)
    story = notes_paras(s)
    if len(story) != 3:
        raise KeyError(f"s46 notes: expected 3 paragraphs, got {len(story)}")
    set_notes_paras(s, story + ["Дальше ХБП без диабета, исследование FIND-CKD."])


# ----------------------------------------------------------------------------- slide 47
def s47(prs):
    s = L.S(prs, 47)
    title(s, "FINEARTS-HF: при сердечной недостаточности с ФВ ЛЖ ≥40% частота ухудшений СН и сердечно-сосудистой смерти "
             "ниже на 16% по сравнению с плацебо", w=660)
    section_tag(s)

    # introduction: two paragraphs, 16 pt
    intro = one(s, "Text_4")
    L.rep(s, "фракцией выброса левого желудочка ≥40%", "ФВ ЛЖ ≥40%")
    L.rep(s, "Доза: до 20 или 40 мг один раз в сутки в зависимости от функции почек. Целевая и максимальная доза по "
             "местной инструкции 20 мг.",
          "Доза в исследовании: до 20 или 40 мг один раз в сутки в зависимости от рСКФ. Максимальная доза "
          "Финеренон-Орвилле: 20 мг в сутки.")
    recolor_text(intro, size=16)
    geom(intro, x=36, y=78, w=888, h=52)
    for p in intro.text_frame.paragraphs:
        p.line_spacing = 0.9
    check("Вводный блок", [para("В анализе: 6001 пациент с симптомной сердечной недостаточностью и ФВ ЛЖ ≥40%, с "
                                "диабетом или без него.", 16),
                           para("Наблюдение: медиана 32 месяца. Доза в исследовании: до 20 или 40 мг один раз в сутки "
                                "в зависимости от рСКФ. Максимальная доза Финеренон-Орвилле: 20 мг в сутки.", 16)],
          888, 52, spacing=0.9)

    # chart caption (one line, 16 pt bold) and chart with patient numbers in the categories
    cap = one(s, "Text_5")
    L.set_lines(cap, ["Суммарное число событий первичного исхода с учётом повторных ухудшений"])
    recolor_text(cap, size=16)
    geom(cap, x=36, y=136, w=530, h=22)
    sh, ch = chart_of(s, "Chart 6")
    set_chart_cats(ch, ["Финеренон (3003 пациента)", "Плацебо (2998 пациентов)"], (1083, 1283), "General")
    geom(sh, x=38, y=164, w=479, h=190)

    # result card on the right: number 32 pt and four lines
    drop(s, "Text_7", "Text_8", "Text_9", "Контроль калия")
    drop(s, "Карточка_результат")
    grey_card(s, 572, 140, 352, 214,
              [para("^^0,84^^", 32, after=2),
               para("Отношение частот (95% ДИ 0,74–0,95); p=0,007", 15, after=1),
               para("Частота событий ниже на 16% по сравнению с плацебо", 15, after=6),
               para("**Ухудшения СН (первые и повторные):** 842 события против 1024, отношение частот 0,82 "
                    "(95% ДИ 0,71–0,94).", 14, after=3),
               para("**Сердечно-сосудистая смерть:** 8,1% на финереноне и 8,7% на плацебо; ОР 0,93 (95% ДИ 0,78–1,11).",
                    14, after=3),
               para("Гиперкалиемия встречалась чаще, чем на плацебо, гипокалиемия реже. Калий контролируют.", 14)],
              "Карточка_результат", anchor="m", pad_y=6)

    # definition of the primary outcome stays as it is, 14 pt
    prim = one(s, "Text_10")
    recolor_text(prim, size=14)
    geom(prim, x=36, y=358, w=888, h=34)
    for p in prim.text_frame.paragraphs:
        p.line_spacing = 0.9

    # two cards at the bottom: FINE-HEART and registration status
    drop(s, "Text_11")
    block(s, 36, 398, 524, 24, 86 - 24 - GAP, TEAL, "FINE-HEART",
          [para("Сводный анализ FIDELIO-DKD, FIGARO-DKD и FINEARTS-HF: 18 991 пациент, медиана 2,9 года. "
                "Госпитализации по поводу СН: ОР 0,83 (95% ДИ 0,75–0,92). Смерть от любой причины: ОР 0,91 "
                "(0,84–0,99). Сердечно-сосудистая смерть: ОР 0,89 (0,78–1,01), p=0,076.", 14)],
          "FINE-HEART", size=16, anchor="m", spacing=0.9)
    block(s, 572, 398, 352, 24, 86 - 24 - GAP, TEAL, "Статус регистрации на 03.10.2026",
          [para("СН с ФВ ЛЖ ≥40%: США (FDA), июль 2025; Евросоюз, 26.03.2026. Япония (22.12.2025) и "
                "Великобритания (апрель 2026): по сообщениям разработчика.", 14)],
          "Статус регистрации", size=16, anchor="m", spacing=0.9)

    footers(s,
            "СН: сердечная недостаточность; ФВ ЛЖ: фракция выброса левого желудочка; рСКФ: расчётная скорость "
            "клубочковой фильтрации; ДИ: доверительный интервал; ОР: отношение рисков.",
            ["Solomon SD, et al. N Engl J Med. 2024;391:1475–1485. Vaduganathan M, et al. Nat Med. 2024;30:3758–3764. "
             "FDA. Supplement Approval Letter, NDA 215341/S-009, July 2025. European Commission Decision C(2026) 2188 "
             "final, 26.03.2026.",
             "Bayer AG. Сообщения о регистрации в Японии (12.2025) и Великобритании (04.2026). Финеренон-Орвилле. "
             "Инструкция по медицинскому применению."])

    # ---- notes
    L.rep(s, "В анализ FINEARTS-HF вошёл 6001 пациент",
          "Первая группа научного раздела: сердечная недостаточность с ФВ ЛЖ 40% и выше. Напомню правило: эту тему "
          "вы сами не начинаете. В анализ FINEARTS-HF вошёл 6001 пациент", notes=True)
    L.rep(s, "Снижение сердечно-сосудистой смерти отдельно не подтверждено: отношение рисков 0,93 с 95% "
             "доверительным интервалом 0,78–1,11.",
          "Ухудшения сердечной недостаточности (первые и повторные): 842 события на финереноне и 1024 на плацебо, "
          "отношение частот 0,82 (95% ДИ 0,71–0,94). Сердечно-сосудистая смерть: 8,1% и 8,7%, отношение рисков 0,93 "
          "(95% ДИ 0,78–1,11), различие статистически незначимо.", notes=True)
    L.rep(s, "встречался чаще при финереноне, чем при плацебо. Поэтому контроль калия остаётся необходимым.",
          "встречался чаще при финереноне, чем при плацебо, а гипокалиемия реже. Контроль калия остаётся "
          "необходимым.", notes=True)
    L.rep(s, "В исследовании максимальная доза составляла 20 или 40 мг один раз в сутки в зависимости от функции "
             "почек. В предоставленной инструкции Финеренон-Орвилле целевая и максимальная доза 20 мг один раз в "
             "сутки, а самостоятельное показание при этой форме сердечной недостаточности не указано. Сопутствующая "
             "СН у пациента с ХБП при диабете 2 типа сама по себе не исключает применение по местному показанию. "
             "Соответствие пациента этому показанию оценивает врач. Зарубежные регистрации и исследовательскую "
             "схему нельзя автоматически переносить в местное применение.",
          "Статусы регистрации при СН с ФВ ЛЖ 40% и выше на 03.10.2026: США (FDA), июль 2025; Евросоюз, решение "
          "Еврокомиссии 26.03.2026 (разработчик объявил 30.03.2026); Япония, 22.12.2025, и Великобритания, апрель "
          "2026, по сообщениям разработчика. Список стран неполный; актуальный список ведёт медицинский отдел. "
          "Рекомендации: по сообщению разработчика, в японских рекомендациях JCS/JHFS 2025 финеренон имеет класс IIa "
          "при СН с ФВ ЛЖ 40% и выше; класс и уровень в рекомендациях ESC 2026 сверяет медицинский отдел.",
          notes=True)
    story = notes_paras(s)
    if len(story) != 4:
        raise KeyError(f"s47 notes: expected 4 paragraphs, got {len(story)}")
    heart = ("Сводный анализ FINE-HEART объединил FIDELIO-DKD, FIGARO-DKD и FINEARTS-HF: 18 991 пациент, медиана "
             "наблюдения 2,9 года. Госпитализации по поводу СН: ОР 0,83 (95% ДИ 0,75–0,92). Смерть от любой причины: "
             "11,0% на финереноне и 12,0% на плацебо, ОР 0,91 (0,84–0,99). Сердечно-сосудистая смерть, первичный "
             "исход анализа: 4,4% и 5,0%, ОР 0,89 (0,78–1,01), p=0,076. Составной почечный исход: ОР 0,80 "
             "(0,72–0,90). Участники FIDELIO-DKD и FIGARO-DKD относятся к группе показания Финеренон-Орвилле.")
    set_notes_paras(s, story[:3] + [heart, story[3],
        "Следующая группа: ХБП при диабете 1 типа.",
        ask("Какая доза?", "в FINEARTS-HF максимальная доза была 20 или 40 мг в зависимости от рСКФ, в инструкции "
            "Финеренон-Орвилле максимальная доза 20 мг один раз в сутки."),
        ask("А пациент с ХБП, диабетом 2 типа и СН?", "такой пациент входит в показание Финеренон-Орвилле, снижение "
            "риска госпитализации по поводу СН включено в показание (FIGARO-DKD, ОР 0,71); соответствие пациента "
            "показанию оценивает врач. Вопросы о СН без ХБП передавайте в медицинский отдел."),
    ])


# ----------------------------------------------------------------------------- run
def run(prs):
    s41(prs)
    s42(prs)
    s43(prs)
    s44(prs)
    s45(prs)
    s46(prs)
    s47(prs)
