"""Last step of the promo rework: junction notes between ranges, footer colours, logo position,
and the new slide order (analysis_promo_47.md, section 3). Runs after impl_01_08 ... impl_41_47 and
impl_new (which appends H1-H6 as slides 48-53)."""
import re

from pptx.dml.color import RGBColor
from pptx.util import Pt

import r14lib as L
from lib import iter_shapes

# old number -> position in the new deck; 48..53 are H1..H6, slide 6 goes to the appendix
ORDER = [1, 48, 2, 3, 4, 5, 7, 8, 9,
         10, 11, 12, 49,
         13, 15,
         16, 17, 18, 50,
         14, 19, 20, 21, 22, 23, 24,
         26, 27, 25,
         31, 32, 33, 28, 29, 30, 34, 35, 36,
         51, 52, 53,
         37, 38, 39, 40, 41,
         42, 47, 45, 46, 44, 43,
         6]

# Repeated transitions: the end of one slide's notes already leads to the next slide, so the
# echo at the start of the next slide's notes goes (old numbers; 52 = H5, 53 = H6).
JUNCTIONS = [
    (19, "Теперь посмотрим, что показали исследования у пациентов.", ""),
    (20, "Сначала почки: исследование FIDELIO-DKD.", ""),
    (21, "Почки мы видели в FIDELIO-DKD. Теперь сердце: исследование FIGARO-DKD, в нём группа пациентов шире, "
         "включая начальные стадии ХБП.",
     "В FIGARO-DKD группа пациентов шире, чем в FIDELIO-DKD, и включает начальные стадии ХБП."),
    (23, "Мы видели исходы. Теперь вернёмся к первому признаку ответа: альбуминурии в первые месяцы лечения.", ""),
    (24, "Альбуминурию снижает и ингибитор НГЛТ2. Часто возникает вопрос, как финеренон сочетается с ним. "
         "На эту тему есть исследование CONFIDENCE.", ""),
    (27, "Мы разобрали, как финеренон сочетается с ингибитором НГЛТ2. Теперь вопрос, который врачи задают часто: "
         "чем финеренон отличается от спиронолактона, который они знают дольше.", ""),
    (25, "Мы разобрали, чем финеренон отличается от спиронолактона. Рекомендации называют именно нестероидный "
         "антагонист МКР с доказанной пользой. Посмотрим, что в них написано. ", ""),
    (31, "Рекомендации определили место финеренона. Теперь практическое правило: как начать лечение. ", ""),
    (32, "Мы разобрали, как начать лечение. Теперь о том, когда проверять анализы. ", ""),
    (33, "Мы выяснили, когда проверять калий и рСКФ. Теперь посмотрим, что делают с дозой по результату анализа. ", ""),
    (28, "Мы разобрали, как начинать лечение, когда проверять анализы и как корректировать дозу. Теперь объясним, "
         "откуда берётся повышение калия и какой оно величины. ", ""),
    (29, "Мы видели, что калий зависит от концентрации финеренона. Её меняют лекарства, которые влияют на фермент "
         "CYP3A4. ", ""),
    (30, "Мы разобрали, как лекарства меняют концентрацию финеренона. Теперь соберём все группы взаимодействий "
         "в одном слайде. ", ""),
    (36, "Теперь другой случай: калий повысился уже на лечении. ", ""),
    (52, "Первая группа возражений касалась других препаратов. Вторая касается калия, функции почек, пациента и "
         "стоимости.", ""),
    (53, "Мы разобрали возражения. Теперь соберём сценарий визита из пяти шагов, и на нём можно потренироваться "
         "в тройках.", "Сценарий из пяти шагов удобно отработать в тройках."),
    (37, "Вернёмся к двум ситуациям из начала занятия. ", ""),
    (42, "Промоционная часть закончена. Дальше научный раздел: данные по другим группам пациентов. Он нужен",
     "Промоционная часть закончена. Научный раздел нужен"),
    (47, "Первая группа научного раздела: сердечная недостаточность с ФВ ЛЖ 40% и выше. ", ""),
    (45, "От сердечной недостаточности переходим к почкам: ХБП при диабете 1 типа. ", ""),
    (44, "Теперь ХБП без диабета. ", ""),
    (43, "Последний слайд раздела: справка к вопросу «чем финеренон отличается от спиронолактона». ",
     "Слайд отвечает на вопрос «чем финеренон отличается от спиронолактона». "),
    # not confirmed (developer's report only): no guideline class in the notes
    (47, " Рекомендации: по сообщению разработчика, в японских рекомендациях JCS/JHFS 2025 финеренон имеет класс "
         "IIa при СН с ФВ ЛЖ 40% и выше; класс и уровень в рекомендациях ESC 2026 сверяет медицинский отдел.", ""),
]


def junctions(prs):
    """Apply JUNCTIONS in the notes; a paragraph that the edit leaves empty is removed."""
    from pptx.oxml.ns import qn
    for n, old, new in JUNCTIONS:
        s = L.S(prs, n)
        body = s.notes_slide.notes_text_frame._txBody
        hit = [p for p in body.findall(qn("a:p"))
               if L._norm(old) in L._norm("".join(t.text or "" for t in p.iter(qn("a:t"))))]
        L.rep(s, old, new, notes=True)
        for p in hit:
            if not "".join(t.text or "" for t in p.iter(qn("a:t"))).strip():
                body.remove(p)


def late(prs):
    """Small fixes decided at the merge (old numbers)."""
    # CREDENCE: 2.6 years is the median follow-up, the share is over the whole trial
    L.rep(L.S(prs, 15), "первичный исход у 11,1% за 2,6 года",
          "первичный исход у 11,1% пациентов (медиана наблюдения 2,6 года)")
    # unconfirmed or too broad wording in the notes (old 27, 25, 32)
    L.rep(L.S(prs, 27), "и только у него есть крупные исследования исходов", "и у него есть крупные исследования исходов",
          notes=True)
    L.rep(L.S(prs, 25), "Ответ: Рекомендации ESC и ERA 2026 по сердечно-сосудистым болезням и ХБП мы сверяем, класс и "
                        "уровень назову после проверки.",
          "Ответ: «Уточню в медицинском отделе и вернусь с ответом».", notes=True)
    L.rep(L.S(prs, 32), "примерно каждые 4 месяца, KDIGO 2022 предлагает такой же интервал.", "примерно каждые 4 месяца.",
          notes=True)
    # sub-headings on the start-of-treatment slide (old 31): purple stays for finerenone only
    from pptx.dml.color import RGBColor as _RGB
    for needle in ("Калий сыворотки", "Начальная доза по функции почек"):
        sh = L.shape_with(L.S(prs, 31), needle)
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.color and r.font.color.type is not None and str(r.font.color.rgb) == "7030A0":
                    r.font.color.rgb = _RGB(0x15, 0x60, 0x82)
    # instruction: at potassium <=5.0 restarting at 10 mg is «considered», not automatic (old 36, 39)
    for n, a, b, notes in (
            (36, "повторяют анализ и возобновляют с 10 мг", "повторяют анализ и рассматривают возобновление с 10 мг", False),
            (36, "и лечение возобновляют с 10 мг", "и врач возобновляет лечение с 10 мг", True),
            (39, "повторяют анализ и возобновляют с 10 мг", "повторяют анализ и рассматривают возобновление с 10 мг", False),
            (39, "лечение возобновляют в дозе 10 мг", "врач может возобновить лечение в дозе 10 мг", False),
            (39, "когда калий 5,0 и ниже, возобновляют с 10 мг", "когда калий 5,0 и ниже, врач может возобновить его с 10 мг",
             True)):
        try:
            L.rep(L.S(prs, n), a, b, notes=notes)
        except KeyError as e:
            print("late: not found", n, e)
    # last repeat on a junction and two language points (old 26, 29, 52 = H5)
    L.rep(L.S(prs, 26), "В CONFIDENCE мы увидели, что финеренон и эмпаглифлозин вместе снижают альбуминурию сильнее, чем "
                        "каждый по отдельности. Теперь посмотрим, как это выглядит в схеме лечения. ", "", notes=True)
    L.rep(L.S(prs, 29), "Концентрация финеренона в крови повышается. Это может увеличивать риск гиперкалиемии.",
          "Концентрация финеренона в крови повышается, и риск гиперкалиемии может расти.", notes=True)
    L.rep(L.S(prs, 52), "не является экономической оценкой.", "не заменяет экономическую оценку.", notes=True)
    # language: no sentence starts with «Это» (text exists only with fix2_09_16 applied)
    try:
        L.rep(L.S(prs, 14), "Это короткая памятка:", "Короткая памятка:", notes=True)
    except KeyError:
        pass


def tidy_notes(prs):
    """No empty paragraphs at the start or end of the notes."""
    from pptx.oxml.ns import qn
    for s in prs.slides:
        if not s.has_notes_slide:
            continue
        body = s.notes_slide.notes_text_frame._txBody
        for end in (0, -1):
            while True:
                ps = body.findall(qn("a:p"))
                if len(ps) < 2:
                    break
                p = ps[end]
                if "".join(t.text or "" for t in p.iter(qn("a:t"))).strip():
                    break
                body.remove(p)


def footers(prs):
    """Sources 8 pt black, abbreviations black (size kept), no hyperlinks."""
    for s in prs.slides:
        for sh in iter_shapes(s.shapes):
            if not sh.has_text_frame:
                continue
            nm = sh.name
            if nm.startswith("Источники") or nm.startswith("Сокращения"):
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        if nm.startswith("Источники"):
                            r.font.size = Pt(8)
                        rPr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                        if rPr is not None:
                            for h in rPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}hlinkClick"):
                                rPr.remove(h)


def logos(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.shape_type == 13 and "огот" in sh.name:
                sh.left, sh.top, sh.width, sh.height = Pt(898), Pt(500), Pt(56), Pt(34)


def reorder(prs):
    n = len(prs.slides)
    assert n == 53, n
    assert sorted(ORDER) == list(range(1, 54)), sorted(ORDER)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for k in ORDER:
        lst.append(ids[k - 1])


def checks(prs):
    bad = []
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in iter_shapes(s.shapes) if sh.has_text_frame]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        for where, t in [("slide", "\n".join(texts)), ("notes", notes)]:
            if "—" in t:
                bad.append((i, where, "em dash"))
            for m in re.finditer(r"[Сс]лайд(?:е|а|ы|ах)?\s+\d+", t):
                bad.append((i, where, m.group(0)))
            if where == "slide" and re.search(r"местн\w* инструкц|предоставленн\w* инструкц|европейск\w* инструкц", t):
                bad.append((i, where, "instruction mention"))
    for b in bad:
        print("CHECK", b)


def run(prs):
    junctions(prs)
    late(prs)
    tidy_notes(prs)
    footers(prs)
    logos(prs)
    reorder(prs)
    checks(prs)
